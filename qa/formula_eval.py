"""Safe stdlib formula evaluator for QA (plan 02-14, closes WR-03).

Re-computes workbook formula strings from their literal inputs, ignoring the
cached values, so a wrong reference inside a formula string is caught.
Test-time guard only: it reads workbook bytes and never saves anything.
It does not import the reference model (checkmodel).

Where Python and Excel semantics differ, the evaluator refuses (raises
FormulaEvalError) rather than return a plausible but wrong value:
- text in arithmetic (Excel coerces "2"*3 to 6 and gives #VALUE! for "a"*3);
- a range used as a scalar operand or as the formula result (Excel would
  apply implicit intersection or spill);
- chained comparisons such as a<b<c (Python chains, Excel evaluates
  (a<b)<c);
- comparisons between different types (number vs text vs logical);
- INDEX on a 2-D range (Excel indexes rows, not a flattened list).
Text comparison is case-insensitive, as in Excel ("LB"="lb" is TRUE).
Errors are lazy values (class _Err): MATCH/INDEX/RANK/IRR misses, division by
zero and #NUM! results are returned as error values that propagate through
arithmetic, comparison and concatenation, are swallowed only by IFERROR (and
ISNUMBER), and are never evaluated away by IF, so IF(guard,IRR(..),"n/a")
behaves as in Excel. A cell whose final value is an error value raises.
Further refusals: unary minus directly before ^ (Excel gives (-2)^2=4,
Python -4), chained ^ (Excel is left-associative), & with a non-integer
number or a logical (locale formatting; whole numbers such as years are
written as General format does, "2035"), ROUND with digits other than 0, COUNTIF
criteria with operators or wildcards, and an IRR whose NPV function does not
have exactly one root on [-0.99, 10] (a non-unique IRR is refused, not guessed).
Text comparison is case-insensitive, as in Excel ("LB"="lb" is TRUE).
"""
import io
import math
import re
from typing import Dict, List, Tuple

import openpyxl
from openpyxl.formula import Tokenizer
from openpyxl.utils import range_boundaries


class FormulaEvalError(ValueError):
    """Unsupported token, unknown function, cycle or Excel error value."""


class _Err(object):
    """Excel error value carried as data (lazy), e.g. #N/A or #NUM!."""
    __slots__ = ("code",)

    def __init__(self, code):
        self.code = code

    def __repr__(self):
        return "_Err(%s)" % self.code


class _Rng(list):
    """Range values in row-major order, with the range shape kept."""

    def __init__(self, vals, rows, cols):
        list.__init__(self, vals)
        self.rows = rows
        self.cols = cols


def _flat(args):
    for x in args:
        if isinstance(x, list):
            for y in _flat(x):
                yield y
        else:
            yield x


def _num(x):
    return isinstance(x, (int, float)) and not isinstance(x, bool)


def _SUM(*a):
    return sum(x for x in _flat(a) if _num(x))


def _AVERAGE(*a):
    v = [x for x in _flat(a) if _num(x)]
    if not v:
        return _Err("#DIV/0!")
    return sum(v) / len(v)


def _MIN(*a):
    v = [x for x in _flat(a) if _num(x)]
    return min(v) if v else 0


def _MAX(*a):
    v = [x for x in _flat(a) if _num(x)]
    return max(v) if v else 0


def _ABS(x):
    return abs(x)


def _AND(*a):
    return all(bool(x) for x in _flat(a) if x is not None)


def _OR(*a):
    vals = [x for x in _flat(a) if x is not None]
    if any(isinstance(x, str) for x in vals):
        raise FormulaEvalError("text in OR is not supported")
    return any(bool(x) for x in vals)


def _NOT(x):
    if isinstance(x, (str, list)):
        raise FormulaEvalError("NOT needs a logical or a number")
    return not x


def _MATCH(x, rng, kind=1):
    vals = list(_flat([rng]))
    if kind == 0:
        for i, b in enumerate(vals):
            if isinstance(x, str) and isinstance(b, str):
                if x.casefold() == b.casefold():
                    return i + 1
            elif _num(x) and _num(b) and x == b:
                return i + 1
            elif isinstance(x, bool) and isinstance(b, bool) and x == b:
                return i + 1
        return _Err("#N/A")
    if kind != 1:
        raise FormulaEvalError("MATCH only supports type 0 and 1")
    pos = None
    for i, b in enumerate(vals):
        if _num(b) and _num(x) and b <= x:
            pos = i + 1
    if pos is None:
        return _Err("#N/A")
    return pos


def _INDEX(rng, k, col=None):
    if col is not None:
        raise FormulaEvalError("INDEX supports 1-D ranges only")
    if not isinstance(rng, list):
        raise FormulaEvalError("INDEX needs a range")
    if isinstance(rng, _Rng) and rng.rows > 1 and rng.cols > 1:
        raise FormulaEvalError("INDEX on a 2-D range is not supported")
    vals = list(_flat([rng]))
    k = int(k)
    if k < 1 or k > len(vals):
        return _Err("#REF!")
    return vals[k - 1]


def _CHOOSE(i, *a):
    i = int(i)
    if i < 1 or i > len(a):
        return _Err("#VALUE!")
    return a[i - 1]


def _IF(c, a, b=False):
    if isinstance(c, _Err):
        return c
    if isinstance(c, (str, list)):
        raise FormulaEvalError("#VALUE! IF condition must be a scalar logical or number")
    return a if c else b


def _IFERROR(a, b):
    return b if isinstance(a, _Err) else a


def _ISNUMBER(x):
    if isinstance(x, list):
        raise FormulaEvalError("ISNUMBER needs a scalar")
    return _num(x)


def _COUNT(*a):
    return sum(1 for x in _flat(a) if _num(x))


def _SUMPRODUCT(*a):
    if not a:
        raise FormulaEvalError("SUMPRODUCT needs arguments")
    shape = None
    cols = []
    for r in a:
        if not isinstance(r, _Rng):
            raise FormulaEvalError("SUMPRODUCT needs ranges")
        if shape is None:
            shape = (r.rows, r.cols)
        elif shape != (r.rows, r.cols):
            raise FormulaEvalError("#VALUE! SUMPRODUCT shape mismatch")
        cols.append([x if _num(x) else 0 for x in r])
    total = 0
    for items in zip(*cols):
        p = 1
        for x in items:
            p *= x
        total += p
    return total


def _NPV(rate, *a):
    if not _num(rate):
        raise FormulaEvalError("NPV rate must be a number")
    if rate == -1:
        return _Err("#DIV/0!")
    return sum(v / (1.0 + rate) ** (i + 1)
               for i, v in enumerate(x for x in _flat(a) if _num(x)))


_IRR_LO, _IRR_HI, _IRR_N = -0.99, 10.0, 6000


def _IRR(values, guess=None):
    flows = [x for x in _flat([values]) if _num(x)]
    if not (any(x > 0 for x in flows) and any(x < 0 for x in flows)):
        return _Err("#NUM!")

    def f(r):
        d = 1.0 + r
        return sum(v / d ** i for i, v in enumerate(flows))

    # uniform grid in ln(1+r): resolves roots near -99% as well as near 1000%
    lo, hi = math.log(1.0 + _IRR_LO), math.log(1.0 + _IRR_HI)
    rs = [math.exp(lo + (hi - lo) * k / _IRR_N) - 1.0 for k in range(_IRR_N + 1)]
    fs = [f(r) for r in rs]
    roots = []
    for k, fk in enumerate(fs):
        if fk == 0:
            roots.append((rs[k], rs[k]))
        elif k + 1 < len(fs) and fs[k + 1] != 0 and fk * fs[k + 1] < 0:
            roots.append((rs[k], rs[k + 1]))
    if len(roots) != 1:
        raise FormulaEvalError(
            "IRR refused: NPV(r) has %d roots on [%s, %s], not exactly one"
            % (len(roots), _IRR_LO, _IRR_HI))
    a, b = roots[0]
    if a == b:
        return a
    fa = f(a)
    for _ in range(200):
        m = (a + b) / 2.0
        fm = f(m)
        if fm == 0 or (b - a) < 1e-13:
            return m
        if (fa < 0) == (fm < 0):
            a, fa = m, fm
        else:
            b = m
    return (a + b) / 2.0


def _MOD(a, b):
    if not (_num(a) and _num(b)):
        raise FormulaEvalError("MOD needs numbers")
    if b == 0:
        return _Err("#DIV/0!")
    return a % b


def _ROUND(x, digits=0):
    if not _num(x):
        raise FormulaEvalError("ROUND needs a number")
    if digits != 0:
        raise FormulaEvalError("ROUND is supported with digits 0 only")
    # half away from zero; round(.., 9) strips binary noise like 0.49999999999999994
    n = int(math.floor(round(abs(float(x)), 9) + 0.5))
    return -n if x < 0 else n


def _RANK(x, ref, order=0):
    if not _num(x):
        return _Err("#VALUE!")
    if order not in (0, 1):
        raise FormulaEvalError("RANK order must be 0 or 1")
    vals = [v for v in _flat([ref]) if _num(v)]
    if x not in vals:
        return _Err("#N/A")
    if order == 1:
        return 1 + sum(1 for v in vals if v < x)
    return 1 + sum(1 for v in vals if v > x)


def _COUNTIF(rng, crit):
    if not isinstance(rng, list):
        raise FormulaEvalError("COUNTIF needs a range")
    if isinstance(crit, list):
        raise FormulaEvalError("COUNTIF criterion must be a scalar")
    vals = list(_flat([rng]))
    if isinstance(crit, str):
        if crit[:1] in ("<", ">", "=") or any(c in crit for c in "*?~"):
            raise FormulaEvalError("COUNTIF supports plain equality criteria only")
        try:
            float(crit)
        except ValueError:
            pass
        else:
            raise FormulaEvalError("COUNTIF numeric-looking text criterion refused")
        return sum(1 for v in vals if isinstance(v, str) and v.casefold() == crit.casefold())
    if isinstance(crit, bool):
        return sum(1 for v in vals if isinstance(v, bool) and v == crit)
    if _num(crit):
        return sum(1 for v in vals if _num(v) and v == crit)
    raise FormulaEvalError("unsupported COUNTIF criterion %r" % (crit,))


def _concat_text(x):
    """Operand of &: text as is; a whole number (such as a year) as Excel's
    General format writes it ("2035"). Other numbers and logicals are
    refused, since their text depends on locale and number format."""
    if isinstance(x, str):
        return x
    if (isinstance(x, (int, float)) and not isinstance(x, bool)
            and math.isfinite(x) and float(x).is_integer() and abs(x) < 1e15):
        return str(int(x))
    raise FormulaEvalError("& is supported between text and whole numbers only")


_FUNCS = {"SUM": _SUM, "AVERAGE": _AVERAGE, "MIN": _MIN, "MAX": _MAX,
          "ABS": _ABS, "AND": _AND, "IF": _IF, "MATCH": _MATCH,
          "INDEX": _INDEX, "CHOOSE": _CHOOSE,
          "SUMPRODUCT": _SUMPRODUCT, "NPV": _NPV, "IRR": _IRR, "MOD": _MOD,
          "ROUND": _ROUND, "RANK": _RANK, "COUNTIF": _COUNTIF, "COUNT": _COUNT,
          "OR": _OR, "NOT": _NOT, "IFERROR": _IFERROR, "ISNUMBER": _ISNUMBER}
# functions that look at error values themselves; all others propagate them
_ERR_AWARE = frozenset(("IF", "IFERROR", "ISNUMBER"))
SUPPORTED_FUNCTIONS = frozenset(_FUNCS)

_INFIX = {"+": "+", "-": "-", "*": "*", "/": "/", "^": "**", "&": "&",
          "=": "==", "<>": "!=",
          "<": "<", ">": ">", "<=": "<=", ">=": ">="}
_COMPARISONS = frozenset(("=", "<>", "<", ">", "<=", ">="))


def _scalar(v, what):
    if isinstance(v, list):
        raise FormulaEvalError("range used as a scalar in %s" % what)
    return v


class _X(object):
    """Typed operand: infix operators follow Excel or refuse.

    Every literal, reference and function result in a translated formula is
    wrapped in _X, so Python's own operator semantics (string repetition,
    list repetition, case-sensitive equality, chaining) are never used.
    """
    __slots__ = ("v",)

    def __init__(self, v):
        self.v = v

    def _n(self, op):
        v = _scalar(self.v, "arithmetic")
        if v is None:
            return 0
        if isinstance(v, bool):
            return int(v)
        if isinstance(v, str):
            raise FormulaEvalError("#VALUE! text %r in arithmetic (%s)" % (v[:20], op))
        if not _num(v):
            raise FormulaEvalError("unsupported value %r in arithmetic" % (v,))
        return v

    def _bad(self, o=None):
        """The first error value among the operands, else None."""
        if isinstance(self.v, _Err):
            return self.v
        if o is not None and isinstance(o.v, _Err):
            return o.v
        return None

    def __add__(self, o):
        e = self._bad(o)
        return _X(e) if e else _X(self._n("+") + o._n("+"))

    def __sub__(self, o):
        e = self._bad(o)
        return _X(e) if e else _X(self._n("-") - o._n("-"))

    def __mul__(self, o):
        e = self._bad(o)
        return _X(e) if e else _X(self._n("*") * o._n("*"))

    def __truediv__(self, o):
        e = self._bad(o)
        if e:
            return _X(e)
        d = o._n("/")
        n = self._n("/")
        if d == 0:
            return _X(_Err("#DIV/0!"))
        return _X(n / d)

    def __pow__(self, o):
        e = self._bad(o)
        if e:
            return _X(e)
        b, x = float(self._n("^")), float(o._n("^"))
        try:
            r = b ** x
        except ZeroDivisionError:
            return _X(_Err("#DIV/0!"))
        except OverflowError:
            return _X(_Err("#NUM!"))
        if isinstance(r, complex):
            return _X(_Err("#NUM!"))
        return _X(r)

    def __and__(self, o):
        e = self._bad(o)
        if e:
            return _X(e)
        a, b = _concat_text(_scalar(self.v, "&")), _concat_text(_scalar(o.v, "&"))
        return _X(a + b)

    def __neg__(self):
        e = self._bad()
        return _X(e) if e else _X(-self._n("unary -"))

    def __pos__(self):
        if isinstance(self.v, _Err):
            return self
        _scalar(self.v, "unary +")
        return self

    def _cmp_pair(self, o):
        a = _scalar(self.v, "comparison")
        b = _scalar(o.v, "comparison")
        # A blank compares as "" against text, 0 against a number and FALSE
        # against a logical, as in Excel.
        if a is None and b is None:
            return 0, 0
        if a is None:
            a = "" if isinstance(b, str) else (False if isinstance(b, bool) else 0)
        if b is None:
            b = "" if isinstance(a, str) else (False if isinstance(a, bool) else 0)
        if isinstance(a, str) and isinstance(b, str):
            return a.casefold(), b.casefold()
        if isinstance(a, bool) and isinstance(b, bool):
            return a, b
        if _num(a) and _num(b):
            return a, b
        raise FormulaEvalError("comparison of different types %r and %r" % (a, b))

    def _cmp(self, o, op):
        e = self._bad(o)
        if e:
            return _X(e)
        a, b = self._cmp_pair(o)
        return _X(op(a, b))

    def __eq__(self, o):
        return self._cmp(o, lambda a, b: a == b)

    def __ne__(self, o):
        return self._cmp(o, lambda a, b: a != b)

    def __lt__(self, o):
        return self._cmp(o, lambda a, b: a < b)

    def __le__(self, o):
        return self._cmp(o, lambda a, b: a <= b)

    def __gt__(self, o):
        return self._cmp(o, lambda a, b: a > b)

    def __ge__(self, o):
        return self._cmp(o, lambda a, b: a >= b)

    __hash__ = None

    def __bool__(self):
        # Python would only need truthiness for a chained comparison.
        raise FormulaEvalError("chained comparison or implicit truth test")


def _unwrap(x):
    return x.v if isinstance(x, _X) else x


def _wrap_func(f, err_aware=False):
    def call(*a):
        args = [_unwrap(x) for x in a]
        if not err_aware:
            for x in args:
                if isinstance(x, _Err):
                    return _X(x)
        return _X(f(*args))
    return call


_SHEET_REF = re.compile(r"^(?:(?P<s>'(?:[^']|'')+'|[A-Za-z0-9_.]+)!)?"
                        r"(?P<r>\$?[A-Za-z]{1,3}\$?[0-9]+(?::\$?[A-Za-z]{1,3}\$?[0-9]+)?)$")
_NAME = re.compile(r"^[A-Za-z_][A-Za-z0-9_]*$")


def _unquote(s):
    if s.startswith("'") and s.endswith("'"):
        return s[1:-1].replace("''", "'")
    return s


class FormulaEvaluator(object):
    def __init__(self, wb_formulas, overrides=None):
        self.wb = wb_formulas
        # keys normalised like cell lookups, so ("Inputs", "$D$182") works
        self.overrides = {(k[0], k[1].replace("$", "").upper()): v
                          for k, v in (overrides or {}).items()}
        self._cache = {}  # type: Dict[Tuple[str, str], object]
        self._busy = set()
        self._names = {}
        for k, v in wb_formulas.defined_names.items():
            self._names[k] = v.attr_text

    def name_target(self, name):
        if name not in self._names:
            raise FormulaEvalError("unknown name %r" % name)
        sheet, ref = self._names[name].rsplit("!", 1)
        return _unquote(sheet), ref.replace("$", "")

    # -- cell and range resolution ------------------------------------
    def value(self, sheet, coord):
        coord = coord.replace("$", "").upper()
        key = (sheet, coord)
        if key in self.overrides:
            return self.overrides[key]
        if key in self._cache:
            return self._cache[key]
        if key in self._busy:
            raise FormulaEvalError("circular reference at %s!%s" % key)
        raw = self.wb[sheet][coord].value
        if isinstance(raw, str) and raw.startswith("="):
            self._busy.add(key)
            try:
                val = self._eval(sheet, raw)
            finally:
                self._busy.discard(key)
        elif raw is None or isinstance(raw, (str, int, float, bool)):
            val = raw
        else:
            raise FormulaEvalError("unsupported cell content at %s!%s: %s"
                                   % (sheet, coord, type(raw).__name__))
        self._cache[key] = val
        return val

    def _ref(self, sheet, ref):
        ref = ref.replace("$", "").upper()
        if ":" not in ref:
            v = self.value(sheet, ref)
            return 0 if v is None else v
        c1, r1, c2, r2 = range_boundaries(ref)
        out = []
        for r in range(r1, r2 + 1):
            for c in range(c1, c2 + 1):
                out.append(self.value(sheet, openpyxl.utils.get_column_letter(c) + str(r)))
        return _Rng(out, r2 - r1 + 1, c2 - c1 + 1)

    # -- translation ---------------------------------------------------
    def _operand(self, sheet, text):
        if "__" in text or "." in text.split("!")[0] and not _SHEET_REF.match(text):
            raise FormulaEvalError("rejected identifier %r" % text)
        m = _SHEET_REF.match(text)
        if m:
            s = _unquote(m.group("s")) if m.group("s") else sheet
            if s not in self.wb.sheetnames:
                raise FormulaEvalError("unknown sheet %r" % s)
            return "_X(_R(%r,%r))" % (s, m.group("r"))
        if _NAME.match(text):
            s, r = self.name_target(text)
            return "_X(_R(%r,%r))" % (s, r)
        raise FormulaEvalError("unsupported operand %r" % text)

    def _translate(self, sheet, formula):
        tok = Tokenizer(formula)
        parts = []
        # comparison operators seen per parenthesis / argument level: a second
        # one at the same level is a chained comparison and is refused
        cmp_depth = [0]
        # per level: prefix sign pending on the current term, and a ^ seen
        pend = [False]
        pows = [0]
        for t in tok.items:
            ty, st, v = t.type, t.subtype, t.value
            if ty == "WHITE-SPACE":
                continue
            if ty == "OPERAND":
                if st == "NUMBER":
                    parts.append("_X(%s)" % (repr(float(v)) if re.search(r"[.eE]", v) else v))
                elif st == "TEXT":
                    parts.append("_X(%r)" % v[1:-1].replace('""', '"'))
                elif st == "LOGICAL":
                    parts.append("_X(True)" if v.upper() == "TRUE" else "_X(False)")
                elif st == "RANGE":
                    parts.append(self._operand(sheet, v))
                else:
                    raise FormulaEvalError("unsupported operand type %s %r" % (st, v))
            elif ty == "FUNC":
                if st == "OPEN":
                    name = v[:-1]
                    if name.upper() != name or name not in SUPPORTED_FUNCTIONS:
                        raise FormulaEvalError("unsupported function %r" % name)
                    parts.append("_F_%s(" % name)
                    cmp_depth.append(0)
                    pend.append(False)
                    pows.append(0)
                else:
                    parts.append(")")
                    cmp_depth.pop()
                    pend.pop()
                    pows.pop()
            elif ty == "SEP":
                if st != "ARG":
                    raise FormulaEvalError("unsupported separator %r" % v)
                parts.append(",")
                cmp_depth[-1] = 0
                pend[-1] = False
                pows[-1] = 0
            elif ty == "PAREN":
                parts.append(v)
                if st == "OPEN":
                    cmp_depth.append(0)
                    pend.append(False)
                    pows.append(0)
                else:
                    cmp_depth.pop()
                    pend.pop()
                    pows.pop()
            elif ty == "OPERATOR-INFIX":
                if v not in _INFIX:
                    raise FormulaEvalError("unsupported operator %r" % v)
                if v in _COMPARISONS:
                    cmp_depth[-1] += 1
                    if cmp_depth[-1] > 1:
                        raise FormulaEvalError("chained comparison in %s" % formula)
                if v == "^":
                    if pend[-1]:
                        raise FormulaEvalError(
                            "unary sign before ^ differs between Excel and Python in %s" % formula)
                    pows[-1] += 1
                    if pows[-1] > 1:
                        raise FormulaEvalError("chained ^ in %s" % formula)
                else:
                    pend[-1] = False
                    pows[-1] = 0
                parts.append(_INFIX[v])
            elif ty == "OPERATOR-PREFIX":
                if v not in ("-", "+"):
                    raise FormulaEvalError("unsupported prefix %r" % v)
                pend[-1] = True
                parts.append(v)
            else:
                raise FormulaEvalError("unsupported token %s %r" % (ty, v))
        return " ".join(parts)

    def _eval(self, sheet, formula):
        py = self._translate(sheet, formula)
        env = {"__builtins__": {}, "_X": _X,
               "_R": lambda s, r: self._ref(s, r)}
        for n, f in _FUNCS.items():
            env["_F_" + n] = _wrap_func(f, n in _ERR_AWARE)
        try:
            out = _unwrap(eval(compile(py, "<formula>", "eval"), env, {}))
            if isinstance(out, _Err):
                if getattr(self, "allow_err", False):
                    return out  # intermediate cell holds an Excel error value
                raise FormulaEvalError("%s in %s" % (out.code, formula))
            return _scalar(out, "the formula result")
        except FormulaEvalError:
            raise
        except ZeroDivisionError:
            raise FormulaEvalError("#DIV/0! in %s" % formula)
        except Exception as e:
            raise FormulaEvalError("cannot evaluate %s: %s" % (formula, e))


def _close(got, want, rel_tol):
    if got == "" and want is None:  # an empty-string result reads back as None
        return True
    if isinstance(got, bool) != isinstance(want, bool):
        return False
    if _num(got) and _num(want):
        return abs(got - want) <= rel_tol * max(1.0, abs(want))
    return got == want


def compare_cached(data, rel_tol=1e-9, overrides=None):
    """Return (mismatches, n_formulas) for the workbook bytes.

    Each formula is evaluated from its string and compared with the cached
    value. Evaluation errors are reported as mismatches, not raised.
    """
    wf = openpyxl.load_workbook(io.BytesIO(data))
    wv = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    ev = FormulaEvaluator(wf, overrides)
    bad = []  # type: List[str]
    n = 0
    for ws in wf:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    n += 1
                    want = wv[ws.title][c.coordinate].value
                    try:
                        got = ev.value(ws.title, c.coordinate)
                    except FormulaEvalError as e:
                        bad.append("%s!%s %s ERROR %s" % (ws.title, c.coordinate, c.value[:80], e))
                        continue
                    if not _close(got, want, rel_tol):
                        bad.append("%s!%s %s got %r want %r" % (
                            ws.title, c.coordinate, c.value[:80], got, want))
    return bad, n


def evaluate_names(data, names, overrides=None, tolerate_errors=False):
    """Evaluate defined names from their formulas under selector overrides.

    A single-cell name returns its scalar value; a range name returns the
    list of cell values in row-major order. The workbook is only read.
    """
    wf = openpyxl.load_workbook(io.BytesIO(data))
    ev = FormulaEvaluator(wf, overrides)
    # tolerate_errors: intermediate cells may hold Excel error values (as Excel
    # does for a pasted invalid selector); IFERROR in the formulas swallows
    # them. A requested name that ends in an error still raises.
    ev.allow_err = tolerate_errors
    out = {}
    for n in names:
        sheet, ref = ev.name_target(n)
        if ":" in ref:
            out[n] = list(ev._ref(sheet, ref))
        else:
            out[n] = ev.value(sheet, ref)
            if isinstance(out[n], _Err):
                raise FormulaEvalError("%s in name %s" % (out[n].code, n))
    return out
