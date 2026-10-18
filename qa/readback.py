"""Read-back verification of the final workbook bytes (QA-01, D-14).

Uses openpyxl for reading only. It never calls save: re-saving through
openpyxl would drop charts and move data validations (threat T-01-29).
"""
import io
from collections import Counter
from typing import Dict, List

import openpyxl
from xlsxwriter.utility import xl_rowcol_to_cell

from qa.lint import lint_formula


class ReadbackError(RuntimeError):
    def __init__(self, problems):
        self.problems = tuple(problems)
        RuntimeError.__init__(self, "read-back failed: " + "; ".join(self.problems))


def _check_inputs(wb, layout, problems):
    # Imported here so qa/ keeps no import-time dependency on xl/.
    from xl.sheets.inputs import INPUT_COLUMNS
    from xl.layout import LABEL_COL
    if "Inputs" not in wb.sheetnames:
        problems.append("Inputs sheet missing")
        return
    ws = wb["Inputs"]
    hdr = layout.header_row + 1
    got = tuple(ws.cell(hdr, LABEL_COL + 1 + i).value
                for i in range(len(INPUT_COLUMNS)))
    if got != INPUT_COLUMNS:
        problems.append("Inputs header differs from the 11 D-10 columns: %r" % (got,))
    for rid, row in layout.record_rows:
        R = row + 1
        for col in "BCEHIJKL":
            if ws["%s%d" % (col, R)].value in (None, ""):
                problems.append("Inputs %s: column %s is empty" % (rid, col))
        status = ws["K%d" % R].value
        if status not in ("Sourced", "Analyst assumption"):
            problems.append("Inputs %s: status %r" % (rid, status))
        elif status == "Analyst assumption":
            if not str(ws["K%d" % R].fill.fgColor.rgb).endswith("FCE4D6"):
                problems.append("Inputs %s: assumption fill missing" % rid)
            if ws["L%d" % R].value in (None, "", "–"):
                problems.append("Inputs %s: assumption has no rationale" % rid)
        elif ws["J%d" % R].value in (None, "", "–"):
            problems.append("Inputs %s: sourced row has no source id" % rid)


def verify_workbook(data, registry, layouts):
    # type: (bytes, tuple, Dict[str, object]) -> List[str]
    problems = []  # type: List[str]
    wb = openpyxl.load_workbook(io.BytesIO(data))
    wbv = openpyxl.load_workbook(io.BytesIO(data), data_only=True)

    found = Counter()
    for ws in wb.worksheets:
        wsv = wbv[ws.title]
        for row in ws.iter_rows():
            for c in row:
                if c.data_type != "f":
                    continue
                found[(ws.title, c.coordinate, c.value)] += 1
                if wsv[c.coordinate].value is None:
                    problems.append("%s!%s: formula has no cached value"
                                    % (ws.title, c.coordinate))
                for p in lint_formula(c.value, ws.title):
                    problems.append("%s!%s: %s" % (ws.title, c.coordinate, p))
    expected = Counter((r.sheet, xl_rowcol_to_cell(r.row, r.col), r.text)
                       for r in registry)
    for key in sorted((expected - found).elements()):
        problems.append("registry formula missing from file: %s!%s %s" % key)
    for key in sorted((found - expected).elements()):
        problems.append("formula in file not in registry (written outside "
                        "the wrapper): %s!%s %s" % key)

    if "inputs" in layouts:
        _check_inputs(wb, layouts["inputs"], problems)

    names = layouts.get("names") or {}
    # Two-way check (CR-01): every user name in the file must have been
    # written through CellWriter.define_name. Built-in _xlnm names (print
    # titles, print area, filters) are sheet settings, not user names.
    file_names = set(n for n in wb.defined_names if not n.startswith("_xlnm."))
    for ws in wb.worksheets:
        file_names |= set("%s!%s" % (ws.title, n) for n in ws.defined_names
                          if not n.startswith("_xlnm."))
    for extra in sorted(file_names - set(names)):
        problems.append("defined name in file not written via CellWriter: %s"
                        % extra)
    for name in sorted(names):
        dn = wb.defined_names.get(name)
        if dn is None:
            problems.append("defined name missing: %s" % name)
        elif dn.attr_text != names[name].lstrip("="):
            problems.append("defined name %s refers to %s, expected %s"
                            % (name, dn.attr_text, names[name]))

    for ws in wb.worksheets:
        if ws.sheet_view.showGridLines is not False:
            problems.append("%s: gridlines are shown" % ws.title)
        if not ws.freeze_panes:
            problems.append("%s: no freeze panes" % ws.title)
    return problems
