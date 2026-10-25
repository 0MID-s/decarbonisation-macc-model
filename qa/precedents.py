"""Transitive formula-dependency walk over a built workbook (plan 04-06).

dependents(data, seeds) returns every formula cell that references a seed
cell directly or through other formulas, resolving defined names (single
cells and ranges) and sheet-qualified references. The workbook is only read:
no formula is evaluated. Functions that make references dynamic (OFFSET,
INDIRECT) are refused, because the walk could not see their targets.
"""
import io
import re
from collections import deque
from typing import Dict, Iterable, List, Set, Tuple

import openpyxl
from openpyxl.formula import Tokenizer
from openpyxl.utils import get_column_letter, range_boundaries

Cell = Tuple[str, str]

_DYNAMIC = ("OFFSET", "INDIRECT", "INDEX_DYNAMIC")
_SHEET_REF = re.compile(
    r"^(?:(?P<s>'(?:[^']|'')+'|[A-Za-z0-9_.]+)!)?"
    r"(?P<r>\$?[A-Za-z]{1,3}\$?[0-9]+(?::\$?[A-Za-z]{1,3}\$?[0-9]+)?)$")
_NAME = re.compile(r"^[A-Za-z_][A-Za-z0-9_]*$")


class PrecedentError(ValueError):
    """A formula the walker cannot analyse safely."""


def _unquote(s):
    if s.startswith("'") and s.endswith("'"):
        return s[1:-1].replace("''", "'")
    return s


def _expand(sheet, ref):
    # type: (str, str) -> List[Cell]
    ref = ref.replace("$", "").upper()
    c1, r1, c2, r2 = range_boundaries(ref)
    return [(sheet, "%s%d" % (get_column_letter(c), r))
            for r in range(r1, r2 + 1) for c in range(c1, c2 + 1)]


def name_cells(data, name):
    # type: (bytes, str) -> List[Cell]
    """Cells a defined name points to (single cell or range)."""
    wb = openpyxl.load_workbook(io.BytesIO(data))
    return _name_cells(wb, name)


def _name_cells(wb, name):
    if name not in wb.defined_names:
        raise PrecedentError("unknown name %r" % name)
    sheet, ref = wb.defined_names[name].attr_text.rsplit("!", 1)
    return _expand(_unquote(sheet), ref)


def _refs(wb, sheet, formula):
    # type: (object, str, str) -> List[Cell]
    out = []  # type: List[Cell]
    for t in Tokenizer(formula).items:
        if t.type == "FUNC" and t.subtype == "OPEN":
            if t.value[:-1].upper() in _DYNAMIC:
                raise PrecedentError("dynamic reference function in %s"
                                     % formula)
        if t.type != "OPERAND" or t.subtype != "RANGE":
            continue
        m = _SHEET_REF.match(t.value)
        if m:
            s = _unquote(m.group("s")) if m.group("s") else sheet
            out.extend(_expand(s, m.group("r")))
        elif _NAME.match(t.value) and t.value in wb.defined_names:
            out.extend(_name_cells(wb, t.value))
        else:
            raise PrecedentError("unresolvable reference %r in %s"
                                 % (t.value, formula))
    return out


def dependents(data, seeds):
    # type: (bytes, Iterable[Cell]) -> Set[Cell]
    wb = openpyxl.load_workbook(io.BytesIO(data))
    rev = {}  # type: Dict[Cell, List[Cell]]
    for ws in wb:
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if isinstance(v, str) and v.startswith("="):
                    me = (ws.title, c.coordinate)
                    for ref in set(_refs(wb, ws.title, v)):
                        rev.setdefault(ref, []).append(me)
    seen = set()  # type: Set[Cell]
    queue = deque((s, c.replace("$", "").upper()) for s, c in seeds)
    while queue:
        cur = queue.popleft()
        for d in rev.get(cur, ()):
            if d not in seen:
                seen.add(d)
                queue.append(d)
    return seen
