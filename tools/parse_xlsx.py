"""Print rows of an .xlsx workbook for human review.

run as .venv/bin/python -W ignore -I tools/parse_xlsx.py PATH [--sheet NAME]
[--match TEXT ...] [--list-sheets]; read-only; prints values for human review;
never writes.
"""
import argparse
import sys

import openpyxl


def main(argv=None):
    ap = argparse.ArgumentParser(description="Print xlsx rows (read-only).")
    ap.add_argument("path")
    ap.add_argument("--sheet")
    ap.add_argument("--match", action="append", default=[])
    ap.add_argument("--list-sheets", action="store_true")
    a = ap.parse_args(argv)
    wb = openpyxl.load_workbook(a.path, read_only=True, data_only=True)
    try:
        if a.list_sheets:
            for n in wb.sheetnames:
                print(n)
            return 0
        needles = [m.lower() for m in a.match]
        for name in wb.sheetnames:
            if a.sheet and name != a.sheet:
                continue
            for r, row in enumerate(wb[name].iter_rows(values_only=True), start=1):
                cells = ["" if c is None else str(c) for c in row]
                text = " ".join(cells).lower()
                if all(n in text for n in needles):
                    if not any(cells):
                        continue
                    print("%s!A%d: %s" % (name, r, " | ".join(cells)))
    finally:
        wb.close()
    return 0


if __name__ == "__main__":
    sys.exit(main())
