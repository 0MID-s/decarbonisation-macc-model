"""Excel desktop recalculation harness (Excel for Mac, driven by osascript).

For each of the 12 selector states it builds the workbook (cached values come
from the reference model), writes a throwaway copy under dist/excel-check-<ts>/,
has Excel set Inputs!G8/G9/G10, recalculate fully and report every formula
cell, then compares numbers at relative tolerance 1e-7 (absolute 1e-9 near
zero). Error values are counted separately. The deliverable is only ever read.
"""
import dataclasses
import datetime
import json
import os
import shutil
import subprocess
import sys

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if ROOT not in sys.path:
    sys.path.insert(0, ROOT)

from checkmodel.exposure_reference import STATES  # noqa: E402

SELECTOR_CELLS = {"scenario": "G8", "basis": "G9", "direct_ets": "G10"}
SELECTOR_NAMES = {"scenario": "Sel_Scenario", "basis": "Sel_Basis",
                  "direct_ets": "Sel_DirectETS"}
JS_PATH = os.path.join(ROOT, "tools", "excel_recalc.js")
JS_CLOSE_PATH = os.path.join(ROOT, "tools", "excel_close.js")
ERROR_MARKERS = ("#DIV/0!", "#N/A", "#NAME?", "#NULL!", "#NUM!", "#REF!",
                 "#VALUE!", "#SPILL!", "#CALC!")
TIMEOUT_S = 300


def make_check_dir(root="dist"):
    """Create dist/excel-check-<UTC timestamp>/; refuse anything outside dist/
    or under data_raw/."""
    base = os.path.realpath(os.path.join(ROOT, "dist"))
    root_abs = root if os.path.isabs(root) else os.path.join(ROOT, root)
    root_abs = os.path.realpath(root_abs)
    parts = root_abs.split(os.sep)
    if "data_raw" in parts:
        raise ValueError("refusing a data_raw path: %s" % root)
    if root_abs != base and not root_abs.startswith(base + os.sep):
        raise ValueError("check folder must be inside dist/: %s" % root)
    stamp = datetime.datetime.now(datetime.timezone.utc).strftime("%Y%m%dT%H%M%S%fZ")
    path = os.path.join(root_abs, "excel-check-" + stamp)
    os.makedirs(path)
    return path


def selector_cells(data):
    """Resolve the three selector cells from the workbook's defined names
    (Sel_Scenario, Sel_Basis, Sel_DirectETS), so the addresses are never
    hard-coded in the JavaScript helper. They must live on Inputs."""
    import io
    import openpyxl
    wf = openpyxl.load_workbook(io.BytesIO(data))
    out = {}
    for key, nm in SELECTOR_NAMES.items():
        sheet, addr = wf.defined_names[nm].attr_text.split("!")
        if sheet.strip("'") != "Inputs":
            raise ValueError("%s is not on the Inputs sheet: %s" % (nm, sheet))
        out[key] = addr.replace("$", "")
    return out


def master_cell(data):
    """Key 'Sheet!A1' of the Chk_Master cell, resolved from the defined name
    (never hard-coded). It must live on the Checks sheet."""
    import io
    import openpyxl
    wf = openpyxl.load_workbook(io.BytesIO(data))
    sheet, addr = wf.defined_names["Chk_Master"].attr_text.split("!")
    sheet = sheet.strip("'")
    if sheet != "Checks":
        raise ValueError("Chk_Master is not on the Checks sheet: %s" % sheet)
    return "%s!%s" % (sheet, addr.replace("$", ""))


def blank_formula_cells(data):
    """Keys of formula cells whose formula can return an empty string; only
    these may legitimately show 0 in Excel against a cached empty value."""
    import io
    import openpyxl
    wf = openpyxl.load_workbook(io.BytesIO(data))
    out = set()
    for ws in wf:
        for row in ws.iter_rows():
            for c in row:
                if (isinstance(c.value, str) and c.value.startswith("=")
                        and '""' in c.value):
                    out.add("%s!%s" % (ws.title, c.coordinate))
    return out


def formula_cells(data):
    """List of (sheet, A1) for every formula cell of the built workbook."""
    import io
    import openpyxl
    wf = openpyxl.load_workbook(io.BytesIO(data))
    out = []
    for ws in wf:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    out.append((ws.title, c.coordinate))
    return out


def cached_values(data):
    """{ 'Sheet!A1': cached value } for every formula cell."""
    import io
    import openpyxl
    wf = openpyxl.load_workbook(io.BytesIO(data))
    wv = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    out = {}
    for ws in wf:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    out["%s!%s" % (ws.title, c.coordinate)] = \
                        wv[ws.title][c.coordinate].value
    return out


def _is_error(v):
    return isinstance(v, str) and v.strip() in ERROR_MARKERS


def compare(cached, excel, rel=1e-7, abs_=1e-9, blank_ok=frozenset()):
    """Return (mismatches, errors). Numbers by tolerance, text by equality;
    an Excel error value is counted in errors (and is not also a mismatch).
    A cached empty value matches an Excel empty value everywhere, but an Excel
    0 only for the keys in blank_ok (formulas that can return an empty text)."""
    mismatches, errors = [], []
    for key, want in cached.items():
        if key not in excel:
            mismatches.append("%s missing from Excel output" % key)
            continue
        got = excel[key]
        if _is_error(got):
            errors.append("%s %s" % (key, got))
            continue
        if isinstance(want, bool) or isinstance(got, bool):
            if got != want:
                mismatches.append("%s got %r want %r" % (key, got, want))
        elif isinstance(want, (int, float)) and isinstance(got, (int, float)):
            if abs(got - want) > max(abs_, rel * max(abs(got), abs(want))):
                mismatches.append("%s got %r want %r" % (key, got, want))
        elif want in (None, "") and (got in (None, "", "missing value")
                                     or (got == 0 and key in blank_ok)):
            continue
        elif got != want:
            mismatches.append("%s got %r want %r" % (key, got, want))
    return mismatches, errors


def _cell_spec(cells):
    from openpyxl.utils.cell import coordinate_from_string, column_index_from_string
    spec = {}
    for sheet, a1 in cells:
        letters, row = coordinate_from_string(a1)
        col = column_index_from_string(letters)
        s = spec.setdefault(sheet, {"maxrow": 1, "maxcol": 1, "cells": []})
        s["maxrow"] = max(s["maxrow"], row)
        s["maxcol"] = max(s["maxcol"], col)
        s["cells"].append([row, col])
    return spec


def _close_in_excel(folder_name):
    """Best effort: close the throwaway workbook in Excel (never saving)
    before its folder is deleted."""
    try:
        subprocess.run(["osascript", "-l", "JavaScript", JS_CLOSE_PATH,
                        folder_name], capture_output=True, text=True,
                       timeout=60)
    except (subprocess.TimeoutExpired, OSError):
        pass


def run_state(state, base=None):
    """Build one state, have Excel recalculate a throwaway copy, compare."""
    import build
    basis, scen, direct = state
    base = base or build.load_dataset("config")
    ds = dataclasses.replace(base, selector_state=(basis, scen),
                             direct_ets_state=direct)
    data, _ = build.build_bytes("main", ds)
    cached = cached_values(data)
    check_dir = make_check_dir()
    try:
        xlsx = os.path.join(check_dir, "check.xlsx")
        cells_json = os.path.join(check_dir, "cells.json")
        out_json = os.path.join(check_dir, "out.json")
        with open(xlsx, "wb") as fh:
            fh.write(data)
        with open(cells_json, "w") as fh:
            json.dump(_cell_spec(formula_cells(data)), fh)
        sel = selector_cells(data)
        cmd = ["osascript", "-l", "JavaScript", JS_PATH, xlsx, cells_json,
               out_json, scen, basis, direct, sel["scenario"], sel["basis"],
               sel["direct_ets"]]
        try:
            proc = subprocess.run(cmd, capture_output=True, text=True,
                                  timeout=TIMEOUT_S)
            if proc.returncode != 0 or not os.path.exists(out_json):
                raise RuntimeError("osascript failed: %s"
                                   % proc.stderr.strip())
        except (subprocess.TimeoutExpired, RuntimeError):
            _close_in_excel(os.path.basename(check_dir))
            raise
        with open(out_json) as fh:
            excel = json.load(fh)
    finally:
        shutil.rmtree(check_dir, ignore_errors=True)
    mism, errs = compare(cached, excel, blank_ok=blank_formula_cells(data))
    master = excel.get(master_cell(data))
    if master != "OK":
        mism.append("Chk_Master in Excel is %r, expected 'OK'" % (master,))
    return len(cached), mism, errs, master


def main():
    import build
    base = build.load_dataset("config")
    print("| Basis | Scenario | Direct ETS | Formula cells | Mismatches | Error values | Chk_Master in Excel |")
    print("|---|---|---|---|---|---|---|")
    bad = False
    details = []
    for state in STATES:
        try:
            n, mism, errs, master = run_state(state, base)
        except (subprocess.TimeoutExpired, RuntimeError, OSError) as exc:
            print("| %s | %s | %s | FAILED | %s | | |" % (
                state[0], state[1], state[2], type(exc).__name__))
            sys.stdout.flush()
            bad = True
            details.append((state, ["state failed: %s" % exc], []))
            continue
        print("| %s | %s | %s | %d | %d | %d | %s |" % (
            state[0], state[1], state[2], n, len(mism), len(errs), master))
        sys.stdout.flush()
        if mism or errs:
            bad = True
            details.append((state, mism[:10], errs[:10]))
    for state, mism, errs in details:
        print("\n%s" % (state,))
        for m in mism + errs:
            print("  " + m)
    return 1 if bad else 0


if __name__ == "__main__":
    sys.exit(main())
