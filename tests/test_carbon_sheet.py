"""Carbon sheet tests (plan 02-13): formulas, cached values, D-44 wording."""
import io
import unittest

import openpyxl

import build
from checkmodel.baseline_reference import model_values
from checkmodel.carbon_reference import (
    SCENARIOS, carbon_paths, eua_gaps, eur_route_factor)
from xl.layout import TIMELINE_FIRST_COL, TIMELINE_LAST_COL


class CarbonSheet(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        ds = build.load_dataset("config")
        data, _ = build.build_bytes("main", ds)
        cls.wb = openpyxl.load_workbook(io.BytesIO(data))
        cls.wbv = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
        cls.ds = ds
        cls.v = model_values(ds.inputs)
        cls.paths = carbon_paths(ds.inputs)

    def _cell(self, wb, name):
        sheet, addr = self.wb.defined_names[name].attr_text.split("!")
        return wb[sheet][addr.replace("$", "").split(":")[0]]

    def _path(self, wb, key):
        sheet, addr = self.wb.defined_names["Path_Carbon_" + key].attr_text.split("!")
        first = addr.replace("$", "").split(":")[0]
        r = wb[sheet][first].row
        return [wb[sheet].cell(row=r, column=c + 1)
                for c in range(TIMELINE_FIRST_COL, TIMELINE_LAST_COL + 1)]

    def test_conversion_steps(self):
        f = self.wb
        self.assertEqual(self._cell(f, "Carb_EUR2010perUSD").value,
                         "=1/Fin_FX_USDperEUR_2010")
        self.assertEqual(self._cell(f, "Carb_HICPRatio2010").value,
                         "=Fin_HICP_EA_2025/Fin_HICP_EA_2010")
        self.assertEqual(self._cell(f, "Carb_USDtoEUR2025").value,
                         "=Carb_EUR2010perUSD*Carb_HICPRatio2010")
        fv = self.wbv
        self.assertAlmostEqual(self._cell(fv, "Carb_EUR2010perUSD").value,
                               1 / 1.3257, delta=1e-4)
        self.assertAlmostEqual(self._cell(fv, "Carb_HICPRatio2010").value,
                               128.75 / 93.14, delta=1e-6)
        self.assertAlmostEqual(self._cell(fv, "Carb_USDtoEUR2025").value,
                               1.0427153041888453, delta=1e-12)
        self.assertAlmostEqual(self._cell(fv, "Carb_USDtoEUR2025").value,
                               eur_route_factor(self.v, 2010), delta=1e-15)

    def test_paths_match_reference(self):
        for key, _t, _n in SCENARIOS:
            cells = self._path(self.wbv, key)
            self.assertEqual(len(cells), 30)
            for c, want in zip(cells, self.paths[key][2]):
                self.assertAlmostEqual(c.value / want, 1.0, delta=1e-12)

    def test_spot_values(self):
        ndc = self._path(self.wbv, "NDC")
        nz = self._path(self.wbv, "NZ50")
        self.assertAlmostEqual(ndc[4].value, 226.17682744007072, delta=1e-9)
        self.assertAlmostEqual(nz[-1].value, 936.6116566345693, delta=1e-9)

    def test_interpolation_formula_text(self):
        c = self._path(self.wb, "NDC")[4]
        self.assertTrue(c.value.startswith("=INDEX($D$"), c.value)
        self.assertIn("-INDEX(", c.value)
        self.assertNotIn("OFFSET", c.value)
        r = c.row
        eff = self.wb["Carbon"].cell(row=r - 2, column=c.column).value
        self.assertTrue(eff.startswith("=MIN("), eff)
        idx = self.wb["Carbon"].cell(row=r - 1, column=c.column).value
        self.assertTrue(idx.startswith("=MATCH("), idx)

    def test_anchor_rows_link_to_inputs(self):
        ws = self.wb["Carbon"]
        found = [c.value for row in ws.iter_rows() for c in row
                 if isinstance(c.value, str) and c.value.startswith("=Inputs!")]
        self.assertEqual(len(found), 36)  # 18 US$ links + 18 point-year links

    def test_anchor_years_fixed_cells(self):
        """CR-01: anchor years never derive from the movable timeline header."""
        ws = self.wb["Carbon"]
        for key, _t, _n in SCENARIOS:
            chk = self._cell(self.wb, "Carb_Check_Anchors_" + key)
            yr_row = None
            for r in range(chk.row - 1, 1, -1):
                if ws.cell(row=r, column=2).value == "NGFS anchor years":
                    yr_row = r
                    break
            self.assertIsNotNone(yr_row)
            n = len(self.paths[key][0])
            for i in range(n):
                f = ws.cell(row=yr_row, column=4 + i).value
                self.assertNotIn("Time_FirstYear", f)
                if i == 0:
                    self.assertEqual(f, "=Carbon_NGFS_Y0Year")
                else:
                    self.assertTrue(f.startswith("=Inputs!"), f)
                    self.assertNotRegex(f, r"\$5$")

    def test_anchor_check_cells(self):
        for key, _t, _n in SCENARIOS:
            self.assertEqual(self._cell(self.wbv, "Carb_Check_Anchors_" + key).value, "OK")
            self.assertIn("IF(AND(", self._cell(self.wb, "Carb_Check_Anchors_" + key).value)

    def test_inputs_point_year_cells(self):
        ws = self.wb["Inputs"]
        for key, _t, _n in SCENARIOS:
            r = None
            for row in ws.iter_rows(min_col=2, max_col=2):
                if row[0].value == "Carbon_NGFS_" + key:
                    r = row[0].row + 2
            self.assertEqual(ws.cell(row=r, column=2).value, "  point year (fixed)")
            ys = [ws.cell(row=r, column=c + 1).value
                  for c in range(TIMELINE_FIRST_COL, TIMELINE_LAST_COL + 1)]
            self.assertEqual([y for y in ys if y is not None],
                             [2030, 2035, 2040, 2045, 2050, 2055])
            # WR-02: point years are locked constants, not unlocked inputs
            for c in range(TIMELINE_FIRST_COL, TIMELINE_LAST_COL + 1):
                cell = ws.cell(row=r, column=c + 1)
                if cell.value is not None:
                    self.assertTrue(cell.protection.locked, cell.coordinate)

    def test_eua_gaps(self):
        """IN-03: read each scenario's own gap cell (column D of its labelled
        row), not any float on the sheet."""
        gaps = eua_gaps(self.ds.inputs)
        wf, wv = self.wb["Carbon"], self.wbv["Carbon"]
        labels = {c.value: c.row for c in wf["B"] if isinstance(c.value, str)}
        f = [c.value for row in wf.iter_rows() for c in row
             if isinstance(c.value, str)
             and c.value.startswith("=Carbon_EUA_Avg2025-")]
        self.assertEqual(len(f), 3)
        for key, tier, _n in SCENARIOS:
            r = labels["Gap: EUA 2025 minus %s NGFS 2025 (EUR 2025)" % tier]
            self.assertTrue(wf["D%d" % r].value.startswith(
                "=Carbon_EUA_Avg2025-$D$"), key)
            self.assertAlmostEqual(wv["D%d" % r].value, gaps[key], delta=1e-9)

    def test_ets2_block(self):
        ws = self.wb["Carbon"]
        forms = [c.value for row in ws.iter_rows() for c in row
                 if isinstance(c.value, str)]
        for want in ("=Carbon_ETS2_StartYear", "=Carbon_ETS2_StabPrice_EUR2025",
                     "=Carbon_ETS2_Scope_Heat", "=Carbon_ETS2_StabRelease"):
            self.assertIn(want, forms)
        wsv = self.wbv["Carbon"]
        vals = [c.value for row in wsv.iter_rows() for c in row]
        self.assertTrue(any(isinstance(x, float) and abs(x - 55.15) < 0.01
                            for x in vals))

    def test_d44_wording_and_limitation(self):
        ws = self.wb["Carbon"]
        txt = " ".join(c.value for row in ws.iter_rows() for c in row
                      if isinstance(c.value, str) and not c.value.startswith("="))
        for want in ("NGFS Phase 5 Scenario Explorer, hosted by IIASA",
                     "https://data.ene.iiasa.ac.at/ngfs/#/license",
                     "Values modified", "2010 real value", "1.31",
                     "indicative analyst estimate"):
            self.assertIn(want, txt)
        self.assertNotIn("CPI-U", txt)

    def test_anchor_limit(self):
        import xl.sheets.carbon as carbon
        self.assertEqual(carbon.GRID_MAX, 9)


if __name__ == "__main__":
    unittest.main()
