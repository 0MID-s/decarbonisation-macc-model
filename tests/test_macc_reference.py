"""Independent MACC reference and helper-table writer tests (plan 01-07)."""
import io
import re
import unittest

import openpyxl
import xlsxwriter

from checkmodel import macc_reference as ref
from config.build_config import BUILD
from tests.fixtures.synthetic import SYNTHETIC_MEASURES, SYNTHETIC_REFERENCE_PRICE
from xl.cells import CellWriter
from xl.macc import MACC_C_BINS, write_macc_a_block, write_macc_c_block
from xl.styles import Styles

OPTIONS = {"in_memory": True, "strings_to_formulas": False,
           "strings_to_urls": False, "strings_to_numbers": False}


class Reference(unittest.TestCase):
    def setUp(self):
        self.sorted = ref.rank_measures(SYNTHETIC_MEASURES)

    def test_rank_sorted_and_stable(self):
        costs = [m[2] for m in self.sorted]
        self.assertEqual(len(costs), 6)
        self.assertEqual(costs, sorted(costs))
        tied = ref.rank_measures([("a", 1, 5.0), ("b", 2, 5.0), ("c", 3, 1.0)])
        self.assertEqual([m[0] for m in tied], ["c", "a", "b"])

    def test_cumulative(self):
        cum = ref.cumulative(self.sorted)
        self.assertEqual(cum[0][0], 0)
        for k in range(5):
            self.assertEqual(cum[k + 1][0], cum[k][1])
        self.assertEqual(cum[-1][1], sum(m[1] for m in SYNTHETIC_MEASURES))

    def test_points(self):
        pts = ref.macc_a_points(self.sorted)
        self.assertEqual(len(pts), 24)
        xs = [p[3] for p in pts]
        self.assertEqual(xs, sorted(xs))
        cum = ref.cumulative(self.sorted)
        for k in range(6):
            blk = pts[4 * k:4 * k + 4]
            self.assertEqual([p[3] for p in blk],
                             [cum[k][0], cum[k][0], cum[k][1], cum[k][1]])
            self.assertEqual([p[2] for p in blk], [0, 1, 1, 0])

    def test_values(self):
        vals = ref.macc_a_values(self.sorted)
        self.assertEqual(len(vals), 24)
        for i, row in enumerate(vals):
            slot, _edge, hflag, _x = ref.macc_a_points(self.sorted)[i]
            self.assertEqual(len(row), 6)
            for s in range(6):
                if s + 1 == slot and hflag:
                    self.assertEqual(row[s], self.sorted[s][2])
                else:
                    self.assertEqual(row[s], 0)
        self.assertTrue(any(v < 0 for row in vals for v in row))

    def test_bins(self):
        bins = ref.macc_c_bins(self.sorted)
        self.assertEqual(len(bins), 200)
        starts = [b[1] for b in bins]
        self.assertEqual(starts[0], 0)
        self.assertTrue(all(b > a for a, b in zip(starts, starts[1:])))
        slots = [b[2] for b in bins]
        self.assertEqual(slots, sorted(slots))
        self.assertEqual((slots[0], slots[-1]), (1, 6))
        total = sum(m[1] for m in self.sorted)
        for k, m in enumerate(self.sorted):
            share = slots.count(k + 1)
            self.assertLessEqual(abs(share - m[1] / total * 200), 1.0)
        vals = ref.macc_c_values(self.sorted)
        self.assertEqual(len(vals), 200)


class HelperTables(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        buf = io.BytesIO()
        wb = xlsxwriter.Workbook(buf, OPTIONS)
        styles = Styles(wb)
        cells = CellWriter(wb)
        wsa = wb.add_worksheet("MACC_A")
        wsc = wb.add_worksheet("MACC_C")
        cls.sorted = ref.rank_measures(SYNTHETIC_MEASURES)
        cls.a = write_macc_a_block(cells, styles, wsa, SYNTHETIC_MEASURES,
                                   SYNTHETIC_REFERENCE_PRICE)
        cls.c = write_macc_c_block(cells, styles, wsc, cls.a, cls.sorted)
        wb.close()
        data = buf.getvalue()
        cls.wb = openpyxl.load_workbook(io.BytesIO(data))
        cls.wbv = openpyxl.load_workbook(io.BytesIO(data), data_only=True)

    def test_constants(self):
        self.assertEqual(MACC_C_BINS, 200)

    def test_ranks_and_sorted(self):
        f, v = self.wb["MACC_A"], self.wbv["MACC_A"]
        self.assertEqual(f["E7"].value, "=RANK(D7,$D$7:$D$12,1)+COUNTIF($D$6:D6,D7)")
        costs = [m[2] for m in SYNTHETIC_MEASURES]
        for i, c in enumerate(costs):
            exp = 1 + sum(1 for x in costs if x < c) + sum(
                1 for x in costs[:i] if x == c)
            self.assertEqual(v["E%d" % (7 + i)].value, exp)
        self.assertEqual(f["C15"].value,
                         "=INDEX($B$7:$B$12,MATCH($B15,$E$7:$E$12,0))")
        self.assertEqual(f["F16"].value, "=SUM($E$14:E15)")
        self.assertEqual(f["G15"].value, "=F15+E15")
        for k, m in enumerate(self.sorted):
            self.assertEqual(v["C%d" % (15 + k)].value, m[0])
            self.assertEqual(v["D%d" % (15 + k)].value, m[2])
            self.assertEqual(v["E%d" % (15 + k)].value, m[1])

    def test_helper_x_and_series(self):
        f, v = self.wb["MACC_A"], self.wbv["MACC_A"]
        self.assertEqual(
            f["E24"].value,
            "=INDEX($F$15:$F$20,$B24)+$C24*INDEX($E$15:$E$20,$B24)")
        self.assertEqual(f["F24"].value,
                         "=IF($B24=F$22,$D24*INDEX($D$15:$D$20,$B24),0)")
        self.assertEqual(f["F23"].value, "=INDEX($C$15:$C$20,F$22)")
        self.assertEqual(f["L24"].value, "=Spike_CarbonPrice")
        pts = ref.macc_a_points(self.sorted)
        vals = ref.macc_a_values(self.sorted)
        for i in range(24):
            self.assertEqual(v["E%d" % (24 + i)].value, pts[i][3])
            for s in range(6):
                self.assertEqual(v.cell(24 + i, 6 + s).value, vals[i][s])
            self.assertEqual(v["L%d" % (24 + i)].value, SYNTHETIC_REFERENCE_PRICE)
        for s in range(6):
            self.assertEqual(v.cell(23, 6 + s).value, self.sorted[s][0])

    def test_macc_c(self):
        f, v = self.wb["MACC_C"], self.wbv["MACC_C"]
        self.assertEqual(f["D5"].value, "=SUM(MACC_A!$E$15:$E$20)")
        self.assertEqual(f["D6"].value, "=COUNT($B$10:$B$209)")
        self.assertEqual(f["D7"].value, "=D5/D6")
        self.assertEqual(f["D10"].value, "=MATCH($C10,MACC_A!$F$15:$F$20,1)")
        self.assertEqual(f["E10"].value,
                         "=IF($D10=E$8,INDEX(MACC_A!$D$15:$D$20,$D10),0)")
        self.assertEqual(v["D6"].value, 200)
        bins = ref.macc_c_bins(self.sorted)
        vals = ref.macc_c_values(self.sorted)
        for i in range(200):
            self.assertEqual(v["C%d" % (10 + i)].value, bins[i][1])
            self.assertEqual(v["D%d" % (10 + i)].value, bins[i][2])
            for s in range(6):
                self.assertEqual(v.cell(10 + i, 5 + s).value, vals[i][s])

    def test_no_literals_in_formulas(self):
        pat = re.compile(r"(?<![A-Za-z$0-9_.])[2-9]\b")
        for ws in self.wb.worksheets:
            for row in ws.iter_rows():
                for c in row:
                    if c.data_type == "f":
                        stripped = re.sub(r"\$?[A-Z]{1,3}\$?[0-9]+", "", c.value)
                        self.assertIsNone(pat.search(stripped), c.value)


if __name__ == "__main__":
    unittest.main()
