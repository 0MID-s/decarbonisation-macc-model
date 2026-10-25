"""Baseline sheet tests: formulas, cached values vs reference, LB/MB rule."""
import io
import unittest

import openpyxl

import build
from qa.lint import LintError
from checkmodel.baseline_reference import compute_baseline, model_values

NAME_TO_FIELD = {
    "BL_FirmVA": "firm_va_meur", "BL_SectorVA": "sector_va_meur",
    "BL_RatioVA": "ratio_va", "BL_RatioTurnover": "ratio_turnover",
    "BL_Gas_MWh": "gas_mwh", "BL_Elec_MWh": "elec_mwh",
    "BL_Heat_MWh": "heat_mwh", "BL_LPG_MWh": "lpg_mwh",
    "BL_Naphtha_MWh": "naphtha_mwh", "BL_Other_MWh": "other_mwh",
    "BL_Total_MWh": "total_mwh", "BL_FuelLoadMW": "fuel_load_mw",
    "BL_ElecBand": "elec_band", "BL_ElecPrice": "elec_price",
    "BL_GasGCVperNCV": "gas_gcv_per_ncv", "BL_GasGJ_GCV": "gas_gj_gcv",
    "BL_GasBand": "gas_band", "BL_GasPriceNCV": "gas_price_ncv",
    "BL_Fleet_MWh": "fleet_mwh", "BL_Fleet_t": "fleet_t",
    "BL_Scope1Stationary": "s1_stationary_t", "BL_Scope1": "s1_total_t",
    "BL_S2_LB": "s2_lb_t", "BL_S2_MB": "s2_mb_t",
    "BL_S12_LB": "s12_lb_t", "BL_S12_MB": "s12_mb_t",
}
RANGE_NAMES = ("Rng_ElecLower", "Rng_ElecS1", "Rng_ElecS2", "Rng_Elec26S1",
               "Rng_GasLower", "Rng_GasS1", "Rng_GasS2", "Rng_Gas26S1")


def _bad_builder(formula):
    def bad(wb, styles, cells):
        ws = wb.get_worksheet_by_name("Baseline")
        cells.formula(ws, 150, 3, formula, styles.get("calc", "tco2e"),
                      value=0.0)
    return bad


class BaselineSheet(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.data, _ = build.build_bytes("main", build.load_dataset("config"))
        cls.wb = openpyxl.load_workbook(io.BytesIO(cls.data))
        cls.wbv = openpyxl.load_workbook(io.BytesIO(cls.data), data_only=True)
        cls.ws = cls.wb["Baseline"]
        cls.wsv = cls.wbv["Baseline"]
        cls.ref = compute_baseline(
            model_values(build.load_dataset("config").inputs))

    def test_check_names(self):
        for nm in ("BL_Check_Carriers", "BL_Check_Ratio"):
            sheet, addr = self._ref(nm).split("!")
            self.assertEqual(sheet, "Baseline")
            self.assertEqual(self.wsv[addr.replace("$", "")].value, "OK")

    def _ref(self, name):
        return self.wb.defined_names[name].attr_text

    def _cell(self, name, wb):
        sheet, addr = self._ref(name).split("!")
        return wb[sheet][addr.replace("$", "")]

    def test_title_and_note(self):
        self.assertIn("reference year 2024", self.ws["B2"].value)
        self.assertIn("net calorific value (NCV)", self.ws["B3"].value)

    def test_names_defined(self):
        for n in tuple(NAME_TO_FIELD) + RANGE_NAMES:
            self.assertIn(n, self.wb.defined_names)

    def test_band_formulas_exact(self):
        self.assertEqual(self._cell("BL_ElecBand", self.wb).value,
                         "=MATCH(BL_Elec_MWh,Rng_ElecLower,1)")
        self.assertEqual(self._cell("BL_GasBand", self.wb).value,
                         "=MATCH(BL_GasGJ_GCV,Rng_GasLower,1)")
        self.assertEqual(self._cell("BL_ElecPrice", self.wb).value[:9],
                         "=AVERAGE(")

    def test_cached_values_match_reference(self):
        for name, field in NAME_TO_FIELD.items():
            got = self._cell(name, self.wbv).value
            want = getattr(self.ref, field)
            self.assertAlmostEqual(got, want, delta=abs(want) * 1e-12,
                                   msg=name)

    def test_every_formula_has_cached_value(self):
        for row in self.ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    self.assertIsNotNone(self.wsv[c.coordinate].value,
                                         c.coordinate)

    def test_scope2_columns_adjacent(self):
        lb = self._ref("BL_S2_LB")
        mb = self._ref("BL_S2_MB")
        self.assertTrue(lb.startswith("Baseline!$D$"))
        self.assertTrue(mb.startswith("Baseline!$E$"))
        self.assertEqual(lb.split("$")[-1], mb.split("$")[-1])

    def test_no_formula_adds_lb_and_mb(self):
        for row in self.ws.iter_rows():
            for c in row:
                v = c.value
                if isinstance(v, str) and v.startswith("="):
                    self.assertFalse("S2_LB" in v and "S2_MB" in v)

    def test_lb_plus_mb_negative(self):
        ds = build.load_dataset("config")
        for f in ("=BL_S2_LB+BL_S2_MB", "=EF_GridLB+EF_ResidualMB",
                  "=BL_S12_LB+BL_S12_MB"):
            with self.assertRaises(LintError, msg=f):
                build.build_bytes("main", ds, (_bad_builder(f),))

    def test_text_statements(self):
        texts = [c.value for row in self.ws.iter_rows() for c in row
                 if isinstance(c.value, str)]
        joined = " ".join(texts)
        for needle in ("Excluded: no official sector-level Scope 3",
                       "dominated by large petrochemical",
                       "Analyst assumption", "lower bound on rated thermal",
                       "memo only: energy balance covers C20 + C21",
                       "never added"):
            self.assertIn(needle, joined)

    def test_checks_ok(self):
        vals = [c.value for row in self.wsv.iter_rows() for c in row
                if c.value in ("OK", "ERROR")]
        self.assertEqual(vals, ["OK", "OK"])


if __name__ == "__main__":
    unittest.main()
