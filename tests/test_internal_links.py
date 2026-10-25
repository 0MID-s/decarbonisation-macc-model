"""Every internal hyperlink points to an existing sheet and a valid cell."""
import io
import re
import unittest

import openpyxl

import build

CELL = re.compile(r"^[A-Z]{1,3}[0-9]{1,7}$")


class InternalLinkTests(unittest.TestCase):
    def _check(self, dataset_name):
        data, _ = build.build_bytes("main", build.load_dataset(dataset_name))
        wb = openpyxl.load_workbook(io.BytesIO(data))
        n = 0
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for c in row:
                    hl = c.hyperlink
                    if hl is None:
                        continue
                    n += 1
                    self.assertIsNone(hl.target, (ws.title, c.coordinate))
                    self.assertTrue(hl.location, (ws.title, c.coordinate))
                    sheet, sep, addr = hl.location.partition("!")
                    self.assertEqual(sep, "!", hl.location)
                    self.assertIn(sheet, wb.sheetnames, hl.location)
                    self.assertRegex(addr, CELL, hl.location)
        return n

    def test_real_workbook_hyperlinks(self):
        self.assertGreater(self._check("config"), 30)

    def test_synthetic_hyperlinks(self):
        self._check("fixtures")


if __name__ == "__main__":
    unittest.main()
