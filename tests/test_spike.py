"""The MACC and dropdown spike target, end to end through build.py."""
import contextlib
import io
import os
import re
import tempfile
import unittest
import zipfile

import openpyxl

import build
from checkmodel import macc_reference as ref
from qa.scan import scan_xlsx_bytes
from tests.fixtures.synthetic import SYNTHETIC_MEASURES


class Spike(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.data, _ = build.build_bytes("spike", build.load_dataset("fixtures"))
        cls.zip = zipfile.ZipFile(io.BytesIO(cls.data))
        cls.wb = openpyxl.load_workbook(io.BytesIO(cls.data))
        cls.wbv = openpyxl.load_workbook(io.BytesIO(cls.data), data_only=True)
        cls.srt = ref.rank_measures(SYNTHETIC_MEASURES)

    def test_two_chart_parts(self):
        charts = [n for n in self.zip.namelist()
                  if re.match(r"xl/charts/chart\d+\.xml$", n)]
        self.assertEqual(len(charts), 2)

    def test_dropdown_sources(self):
        idx = self.wb.sheetnames.index("Dropdown") + 1
        xml = self.zip.read("xl/worksheets/sheet%d.xml" % idx).decode("utf-8")
        sources = re.findall(r"<formula1>([^<]*)</formula1>", xml)
        self.assertEqual(sources, ["Lists!$B$6:$B$8", "List_Scenario",
                                   "List_Basis"])

    def test_cached_values(self):
        for ws in self.wb.worksheets:
            for row in ws.iter_rows():
                for c in row:
                    if c.data_type == "f":
                        self.assertIsNotNone(
                            self.wbv[ws.title][c.coordinate].value, c.coordinate)
        a = self.wbv["MACC_A"]
        pts = ref.macc_a_points(self.srt)
        vals = ref.macc_a_values(self.srt)
        for i in range(24):
            self.assertEqual(a["E%d" % (24 + i)].value, pts[i][3])
            for s in range(6):
                self.assertEqual(a.cell(24 + i, 6 + s).value, vals[i][s])
        c = self.wbv["MACC_C"]
        bins = ref.macc_c_bins(self.srt)
        for i in range(200):
            self.assertEqual(c["D%d" % (10 + i)].value, bins[i][2])
        d = self.wbv["Dropdown"]
        self.assertEqual(d["C10"].value, "Selected (direct): Central")
        self.assertEqual(d["C12"].value, 2)

    def test_scan_clean(self):
        self.assertEqual(scan_xlsx_bytes(self.data), [])

    def test_config_dataset_refused(self):
        with tempfile.TemporaryDirectory() as tmp:
            out = os.path.join(tmp, "s.xlsx")
            with contextlib.redirect_stdout(io.StringIO()), \
                    contextlib.redirect_stderr(io.StringIO()):
                rc = build.main(["--target", "spike", "--dataset", "config",
                                 "--out", out])
            self.assertEqual(rc, 2)
            self.assertFalse(os.path.exists(out))

    def test_instructions_first_with_twelve_checks(self):
        self.assertEqual(self.wb.sheetnames[0], "Instructions")
        ws = self.wb["Instructions"]
        checks = [r[0].value for r in ws.iter_rows(min_col=2, max_col=2)
                  if isinstance(r[0].value, str)
                  and re.match(r"Check \d+$", r[0].value)]
        self.assertEqual(len(checks), 12)
        self.assertIn("GATE-A", ws["B1"].value)


if __name__ == "__main__":
    unittest.main()
