"""Selector band, tier names, fleet names and boolean check format (03-03)."""
import dataclasses
import io
import unittest

import openpyxl
import xlsxwriter

import build
from checkmodel.baseline_reference import compute_baseline, model_values
from xl.styles import Styles, apply_bool_check_format


class Selectors(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        ds = build.load_dataset("config")
        cls.data, _ = build.build_bytes("main", ds)
        cls.wb = openpyxl.load_workbook(io.BytesIO(cls.data))
        cls.wbv = openpyxl.load_workbook(io.BytesIO(cls.data), data_only=True)
        cls.ref = compute_baseline(model_values(ds.inputs))

    def test_selector_names_and_values(self):
        self.assertEqual(self.wb.defined_names["Sel_Scenario"].attr_text,
                         "Inputs!$G$8")
        self.assertEqual(self.wb.defined_names["Sel_Basis"].attr_text,
                         "Inputs!$G$9")
        ws = self.wb["Inputs"]
        self.assertEqual(ws["G8"].value, "Central")
        self.assertEqual(ws["G9"].value, "Market-based")

    def test_validations(self):
        dv = dict((str(d.sqref), d.formula1)
                  for d in self.wb["Inputs"].data_validations.dataValidation)
        self.assertEqual(dv["G8"].lstrip("="), "List_Scenario")
        self.assertEqual(dv["G9"].lstrip("="), "List_Basis")

    def test_header_row(self):
        self.assertEqual(self.wb["Inputs"]["B11"].value, "ID")
        self.assertEqual(self.wb["Inputs"]["N11"].value, "=Time_FirstYear")

    def test_tier_names(self):
        ws = self.wb["Lists"]
        for n, t in (("Tier_First", "Do first"), ("Tier_Next", "Next"),
                     ("Tier_Later", "Later")):
            sheet, addr = self.wb.defined_names[n].attr_text.split("!")
            self.assertEqual(self.wb[sheet][addr.replace("$", "")].value, t)
        self.assertEqual(self.wb.defined_names["List_Tier"].attr_text,
                         "Lists!$F$6:$F$8")

    def test_fleet_names(self):
        for n, want in (("BL_Fleet_km", self.ref.fleet_km),
                        ("BL_Fleet_Litres", self.ref.fleet_litres)):
            sheet, addr = self.wb.defined_names[n].attr_text.split("!")
            got = self.wbv[sheet][addr.replace("$", "")].value
            self.assertAlmostEqual(got, want, places=6)

    def test_bool_format(self):
        buf = io.BytesIO()
        wb = xlsxwriter.Workbook(buf, {"in_memory": True})
        ws = wb.add_worksheet("T")
        apply_bool_check_format(ws, Styles(wb), 0, 0, 3, 0)
        wb.close()
        w = openpyxl.load_workbook(io.BytesIO(buf.getvalue()))["T"]
        rules = [r for rs in w.conditional_formatting for r in rs.rules]
        self.assertEqual(sorted(r.formula[0] for r in rules),
                         ["FALSE", "TRUE"])


class DirectEtsToggle(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.ds = build.load_dataset("config")
        data, _ = build.build_bytes("main", cls.ds)
        cls.wb = openpyxl.load_workbook(io.BytesIO(data))

    def test_default_off_and_names(self):
        self.assertEqual(self.wb["Inputs"]["G10"].value, "Off")
        self.assertEqual(self.wb.defined_names["Sel_DirectETS"].attr_text,
                         "Inputs!$G$10")
        self.assertEqual(self.wb.defined_names["List_OnOff"].attr_text,
                         "Lists!$H$6:$H$7")
        ws = self.wb["Lists"]
        self.assertEqual((ws["H6"].value, ws["H7"].value), ("Off", "On"))

    def test_validation(self):
        dv = dict((str(d.sqref), d.formula1)
                  for d in self.wb["Inputs"].data_validations.dataValidation)
        self.assertEqual(dv["G10"].lstrip("="), "List_OnOff")

    def test_rationale_text(self):
        self.assertIn("20 MW", self.wb["Inputs"]["C10"].value)

    def test_on_state_build(self):
        ds = dataclasses.replace(self.ds, direct_ets_state="On")
        data, _ = build.build_bytes("main", ds)
        ws = openpyxl.load_workbook(io.BytesIO(data))["Inputs"]
        self.assertEqual(ws["G10"].value, "On")
        self.assertEqual(ws["G8"].value, "Central")
        self.assertEqual(ws["G9"].value, "Market-based")

    def test_helper_and_mirror(self):
        from checkmodel import measures_reference as mr
        from xl.sheets import lists
        from xl.selector import direct_ets_state
        self.assertEqual(direct_ets_state(self.ds), "Off")
        self.assertEqual(direct_ets_state(
            dataclasses.replace(self.ds, direct_ets_state="On")), "On")
        self.assertEqual(mr.ONOFF, lists.ONOFF)
        self.assertEqual(mr.DEFAULT_DIRECT_ETS, "Off")


if __name__ == "__main__":
    unittest.main()
