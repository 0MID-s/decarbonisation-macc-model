"""Dashboard tests (plan 05-04). Task 1: lint pre-tests for the non-link
formulas the Dashboard and Checks sheets use."""
import io
import re
import unittest
import zipfile

import openpyxl

import build
from checkmodel.outputs_reference import compute_outputs
from config.inputs import INPUTS
from qa.formula_eval import compare_cached, evaluate_names
from qa.lint import lint_formula

# Recorded outcome (Task 1): the OR of two comparisons against LB- and
# MB-tagged names is ACCEPTED by the lint (pinned below).
OR_FORMULA = "=OR(Dash_AbatedSel=M_Total_tLB,Dash_AbatedSel=M_Total_tMB)"


class LintPreTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        holder = []
        ds = build.load_dataset("config")
        build.build_bytes("main", ds,
                          extra_builders=[lambda wb, s, c: holder.append(c)])
        cls.tags = holder[0].tags

    def ok(self, f):
        self.assertEqual(lint_formula(f, "Dashboard", self.tags), [], f)

    def test_choose_selected_basis(self):
        self.ok("=CHOOSE(M_BasisIdx,M_Total_tLB,M_Total_tMB)")
        self.ok("=CHOOSE(M_BasisIdx,BL_S12_LB,BL_S12_MB)")

    def test_share_and_tile_vs_macc(self):
        self.ok("=Dash_AbatedSel/Dash_S12Sel")
        self.ok("=ABS(Dash_AbatedSel-MACC_TotalAbatement)<="
                "COUNT(MACC_Chart_Width)*Const_MACCRoundTol")

    def test_single_comparisons_pass(self):
        self.ok("=IFERROR(Dash_AbatedSel=M_Total_tLB,FALSE)")
        self.ok("=IFERROR(Dash_AbatedSel=M_Total_tMB,FALSE)")

    def test_or_formula_outcome_recorded(self):
        problems = lint_formula(OR_FORMULA, "Checks", self.tags)
        # Pinned: ACCEPTED, because Dash_AbatedSel is an untagged helper and
        # each operand of the OR is a single comparison with one tagged name.
        self.assertEqual(problems, [])


# selector names, plus the two year inputs of the exposure-window label (CR-02)
SEL = r"(Sel_Scenario|Sel_Basis|Sel_DirectETS|Time_FirstYear|Expo_LastYear)"
STR = r'"[^"]*"'
SHAPES = (
    r"^=[A-Za-z_][A-Za-z0-9_]*$",
    r"^=CHOOSE\(M_BasisIdx,[A-Za-z0-9_]+,[A-Za-z0-9_]+\)$",
    r"^=Dash_AbatedSel/Dash_S12Sel$",
    r"^=(%s|%s)(&(%s|%s))*$" % (STR, SEL, STR, SEL),
)


class DashboardSheet(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.ds = build.load_dataset("config")
        cls.data, _ = build.build_bytes("main", cls.ds)
        cls.wb = openpyxl.load_workbook(io.BytesIO(cls.data))
        cls.ws = cls.wb["Dashboard"]
        cls.ref = compute_outputs(INPUTS, "Market-based", "Central", "Off")

    def _addr(self, name):
        sheet, addr = self.wb.defined_names[name].attr_text.split("!")
        return sheet, addr.replace("$", "")

    def test_links_only_shapes(self):
        n = 0
        for row in self.ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    n += 1
                    self.assertTrue(
                        any(re.match(p, c.value) for p in SHAPES),
                        "%s %s" % (c.coordinate, c.value))
        self.assertGreater(n, 20)

    def _dashboard_charts(self):
        with zipfile.ZipFile(io.BytesIO(self.data)) as z:
            wbx = z.read("xl/workbook.xml").decode()
            rid = re.search(r'<sheet name="Dashboard"[^>]*r:id="([^"]+)"',
                            wbx).group(1)
            rels = z.read("xl/_rels/workbook.xml.rels").decode()
            target = re.search(r'<Relationship [^>]*Id="%s"[^>]*Target="([^"]+)"'
                               % rid, rels) or re.search(
                r'<Relationship [^>]*Target="([^"]+)"[^>]*Id="%s"' % rid, rels)
            sheet_file = target.group(1).split("/")[-1]
            srels = z.read("xl/worksheets/_rels/%s.rels" % sheet_file).decode()
            drawing = re.search(r'Target="\.\./drawings/([^"]+)"', srels).group(1)
            drels = z.read("xl/drawings/_rels/%s.rels" % drawing).decode()
            charts = re.findall(r'Target="\.\./charts/([^"]+)"', drels)
            return [z.read("xl/charts/" + c).decode() for c in charts]

    def test_three_charts_linked_titles(self):
        charts = self._dashboard_charts()
        self.assertEqual(len(charts), 3)
        for x in charts:
            self.assertRegex(
                x, r"<c:title><c:tx><c:strRef><c:f>[^<]+</c:f>")
        self.assertTrue(any("<c:areaChart>" in x for x in charts))
        self.assertTrue(any('<c:grouping val="stacked"/>' in x for x in charts))
        self.assertTrue(any('<c:grouping val="clustered"/>' in x
                            for x in charts))

    def test_print_setup(self):
        ws = self.ws
        self.assertTrue(ws.print_area)
        self.assertTrue(ws.sheet_properties.pageSetUpPr.fitToPage)
        # XlsxWriter omits the OOXML default of 1 for both attributes
        self.assertIn(ws.page_setup.fitToWidth, (None, 1))
        self.assertIn(ws.page_setup.fitToHeight, (None, 1))
        self.assertEqual(ws.page_setup.orientation, "landscape")
        self.assertEqual(int(ws.page_setup.paperSize), 9)

    def test_no_validation_gridlines_panes(self):
        ws = self.ws
        self.assertEqual(len(ws.data_validations.dataValidation), 0)
        self.assertFalse(ws.sheet_view.showGridLines)
        self.assertIsNotNone(ws.freeze_panes)

    def test_cached_values_equal_reference(self):
        wbv = openpyxl.load_workbook(io.BytesIO(self.data), data_only=True)
        r = self.ref

        def val(name):
            sheet, addr = self._addr(name)
            return wbv[sheet][addr].value
        for name, want in (("Dash_NPVFin", r.npv_fin),
                           ("Dash_NPVIncl", r.npv_incl),
                           ("Dash_AbatedSel", r.abated_sel),
                           ("Dash_S12Sel", r.s12_sel),
                           ("Dash_Share", r.share),
                           ("Dash_CostNotActing", r.cum_bau_sel),
                           ("Dash_CostAvoided", r.cum_avoided_sel),
                           ("Dash_CostNotActingDisc", r.disc_bau_sel)):
            self.assertAlmostEqual(val(name), want, delta=1e-6 * max(1, abs(want)),
                                   msg=name)
        self.assertEqual(val("Dash_Flag"), "OK")
        self.assertEqual(val("Dash_FailCount"), 0)
        self.assertEqual(val("Dash_ScenChartTitle"), r.labels["scen_chart"])
        bad, _ = compare_cached(self.data)
        self.assertEqual(bad, [])

    def test_cost_label_and_undiscounted_wording(self):
        text = [c.value for row in self.ws.iter_rows() for c in row
                if isinstance(c.value, str)]
        self.assertTrue(any("undiscounted" in t for t in text))
        self.assertTrue(any("can be negative" in t for t in text))

    def test_cost_window_label_is_live(self):
        # CR-02: the cumulative tiles cover the exposure window only, so the
        # label names it from Time_FirstYear and Expo_LastYear.
        from qa.formula_eval import FormulaEvaluator
        hits = [c for row in self.ws.iter_rows() for c in row
                if isinstance(c.value, str) and "Expo_LastYear" in c.value]
        self.assertEqual(len(hits), 1)
        coord = hits[0].coordinate
        text = [c.value for row in self.ws.iter_rows() for c in row
                if isinstance(c.value, str)]
        self.assertFalse(any("last model year" in t for t in text))
        r = compute_outputs(INPUTS, "Market-based", "Central", "Off")
        self.assertEqual(r.labels["cost_window"], "Cash carbon cost 2026 to "
                         "2035 (exposure window), undiscounted, real EUR 2025")
        wf = openpyxl.load_workbook(io.BytesIO(self.data))
        self.assertEqual(FormulaEvaluator(wf).value("Dashboard", coord),
                         r.labels["cost_window"])
        sheet, addr = self._addr("Expo_LastYear")
        got = FormulaEvaluator(wf, {(sheet, addr): 2040}).value(
            "Dashboard", coord)
        self.assertIn("2026 to 2040 (exposure window)", got)

    def test_low_scenario_negative_cost_master_ok(self):
        sheet, addr = self._addr("Sel_Scenario")
        got = evaluate_names(self.data, ["Dash_CostNotActing", "Chk_Master"],
                             overrides={(sheet, addr): "Low"})
        self.assertLess(got["Dash_CostNotActing"], 0)
        self.assertEqual(got["Chk_Master"], "OK")


if __name__ == "__main__":
    unittest.main()
