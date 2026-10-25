"""Exposure sheet part 1 tests (plan 04-05): channel rows vs the reference."""
import io
import unittest

import openpyxl

import build
from checkmodel.exposure_reference import compute_exposure
from checkmodel.measures_reference import DEFAULT_DIRECT_ETS, DEFAULT_STATE
from qa.formula_eval import compare_cached, evaluate_names
from xl.layout import TIMELINE_FIRST_COL, TIMELINE_LAST_COL

KEYS = ("CP", "NDC", "NZ50")
REL = 1e-9


class ExposureSheet(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        ds = build.load_dataset("config")
        cls.ds = ds
        cls.data, _ = build.build_bytes("main", ds)
        cls.wb = openpyxl.load_workbook(io.BytesIO(cls.data))
        cls.wbv = openpyxl.load_workbook(io.BytesIO(cls.data), data_only=True)
        cls.ref = compute_exposure(ds.inputs, DEFAULT_STATE[0],
                                   DEFAULT_STATE[1], DEFAULT_DIRECT_ETS)

    def _row(self, wb, name):
        sheet, addr = self.wb.defined_names[name].attr_text.split("!")
        r = wb[sheet][addr.replace("$", "").split(":")[0]].row
        return [wb[sheet].cell(row=r, column=c + 1).value
                for c in range(TIMELINE_FIRST_COL, TIMELINE_LAST_COL + 1)]

    def _close(self, name, want):
        got = self._row(self.wbv, name)
        self.assertEqual(len(got), len(want), name)
        for g, x in zip(got, want):
            self.assertLessEqual(abs(g - x), REL * max(1.0, abs(x)), name)

    def test_names_and_cached_values(self):
        r = self.ref
        self._close("Expo_BAU_tStat", r.t_stat_bau)
        self._close("Expo_With_tStat", r.t_stat_with)
        self._close("Expo_With_tFleet", r.t_fleet_with)
        self._close("Expo_With_ElecMWh", r.elec_with)
        for key in KEYS:
            s = r.scen[key]
            self._close("Expo_PETS1_" + key, s.ets1)
            self._close("Expo_PETS2_" + key, s.ets2)
            self._close("Expo_PStat1_" + key, s.p_stat1)
            self._close("Expo_PStat2_" + key, s.p_stat2)
            for case, parts in (("BAU", (s.bau_ets2, s.bau_dir, s.bau_pt,
                                         s.bau_total)),
                                ("With", (s.with_ets2, s.with_dir, s.with_pt,
                                          s.with_total))):
                for nm, vals in zip(("ETS2", "DirETS", "PT", "Total"), parts):
                    self._close("Expo_%s_%s_%s" % (case, nm, key), vals)
            self._close("Expo_Avoided_" + key, s.avoided)

    def test_compare_cached_no_mismatch(self):
        bad, n = compare_cached(self.data)
        self.assertEqual(bad, [], bad[:5])
        self.assertGreater(n, 0)

    def test_ets2_zero_before_start(self):
        for key in KEYS:
            for case in ("BAU", "With"):
                vals = self._row(self.wbv, "Expo_%s_ETS2_%s" % (case, key))
                self.assertEqual(vals[0], 0)
                self.assertEqual(vals[1], 0)
            self.assertEqual(self._row(self.wbv, "Expo_PETS2_" + key)[:2],
                             [0, 0])

    def test_toggle_on(self):
        sheet, ref_ = self.wb.defined_names["Sel_DirectETS"].attr_text.split("!")
        names = ["Expo_PStat2_NDC", "Expo_PStat1_NDC", "Path_ETS1_NDC",
                 "Expo_BAU_Total_NDC", "Expo_With_Total_NDC",
                 "Expo_BAU_ETS2_NDC", "Expo_BAU_DirETS_NDC"]
        out = evaluate_names(self.data, names,
                             overrides={(sheet, ref_): "On"})
        self.assertTrue(all(x == 0 for x in out["Expo_PStat2_NDC"]))
        for g, w in zip(out["Expo_PStat1_NDC"], out["Path_ETS1_NDC"]):
            self.assertAlmostEqual(g, w, delta=1e-9)  # free allocation 0
        on = compute_exposure(self.ds.inputs, DEFAULT_STATE[0],
                              DEFAULT_STATE[1], "On").scen["NDC"]
        for nm, want in (("Expo_BAU_Total_NDC", on.bau_total),
                         ("Expo_With_Total_NDC", on.with_total),
                         ("Expo_BAU_DirETS_NDC", on.bau_dir)):
            for g, x in zip(out[nm], want):
                self.assertLessEqual(abs(g - x), REL * max(1.0, abs(x)), nm)
        # ETS2 on fleet diesel remains, stationary part is gone
        for g, x in zip(out["Expo_BAU_ETS2_NDC"], on.bau_ets2):
            self.assertLessEqual(abs(g - x), REL * max(1.0, abs(x)))

    def test_notes_present(self):
        txt = " ".join(c.value for row in self.wb["Exposure"].iter_rows()
                       for c in row if isinstance(c.value, str)
                       and not c.value.startswith("="))
        self.assertIn("is negative", txt)
        self.assertIn("Never added to the internal carbon price", txt)


if __name__ == "__main__":
    unittest.main()
