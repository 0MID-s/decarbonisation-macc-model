"""Deliverable hardening (D-23, D-24, OQ-3). Deliberate-break and
invalid-selector cases live in tests/test_checks_sheet.py and are not
duplicated here."""
import io
import re
import unittest
import zipfile

import openpyxl

import build
from config.build_config import BUILD
from qa.readback import verify_workbook  # noqa: F401  (build runs it)
from qa.scan import scan_or_raise
from xl.sheets.common import SYNTHETIC_TEXT

D23 = ["Cover", "ReadMe", "Dashboard", "Inputs", "Sources", "Baseline",
       "Rates", "Carbon", "Measures", "MACC", "Exposure", "Sensitivity",
       "Disclosure", "Checks", "Notes", "Lists"]
CACHED_SHEETS = ("Dashboard", "Checks", "Disclosure", "Cover", "ReadMe")


class Deliverable(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        ds = build.load_dataset("config")
        cls.data, _ = build.build_bytes("main", ds)
        cls.zf = zipfile.ZipFile(io.BytesIO(cls.data))
        cls.wb = openpyxl.load_workbook(io.BytesIO(cls.data))
        cls.wbv = openpyxl.load_workbook(io.BytesIO(cls.data),
                                         data_only=True)

    def test_no_assistant_or_session_wording(self):
        """Client-visible text reads as an analyst's model: no "the user"
        and no assistant name in any text-bearing part."""
        import html
        bad = [r"\bthe user\b", r"(?i)claude"]
        hits = []
        for n in self.zf.namelist():
            if not n.endswith(".xml") or n.startswith("xl/styles"):
                continue
            raw = self.zf.read(n).decode("utf-8")
            if n.startswith("xl/worksheets/"):
                parts = re.findall(r"<is>(.*?)</is>", raw, re.S)
                parts += re.findall(r'<c [^>]*t="str"[^>]*>.*?<v>(.*?)</v>',
                                    raw, re.S)
                raw = " ".join(parts)
            text = html.unescape(re.sub(r"<[^>]+>", " ", raw))
            for pat in bad:
                for m in re.finditer(pat, text):
                    hits.append((n, m.group(0), text[max(0, m.start() - 40):
                                                      m.end() + 40]))
        self.assertEqual(hits, [])

    def test_sheet_order(self):
        self.assertEqual(self.wb.sheetnames, D23)

    def test_opens_on_cover_only_selected(self):
        wbxml = self.zf.read("xl/workbook.xml").decode("utf-8")
        m = re.search(r'activeTab="(\d+)"', wbxml)
        self.assertTrue(m is None or m.group(1) == "0")
        sel = [n for n in self.zf.namelist()
               if re.match(r"xl/worksheets/sheet\d+\.xml$", n)
               and 'tabSelected="1"' in self.zf.read(n).decode("utf-8")]
        self.assertEqual(sel, ["xl/worksheets/sheet1.xml"])

    def test_no_spike_or_synthetic(self):
        for n in self.wb.sheetnames:
            self.assertNotIn("spike", n.lower())
            self.assertNotIn("synthetic", n.lower())
        ss = self.zf.read("xl/sharedStrings.xml").decode("utf-8")
        self.assertNotIn(SYNTHETIC_TEXT, ss)
        self.assertNotIn("SYNTHETIC FIXTURE", ss)

    def test_properties(self):
        core = self.zf.read("docProps/core.xml").decode("utf-8")
        app = self.zf.read("docProps/app.xml").decode("utf-8")
        self.assertIn("<dc:creator>0MID-s</dc:creator>", core)
        self.assertIn(BUILD.title, core.replace("&amp;", "&"))
        self.assertIn(BUILD.subject, core.replace("&amp;", "&"))
        self.assertNotIn("@", core)
        self.assertNotIn("@", app)

    def test_cached_values_present(self):
        missing = []
        for name in CACHED_SHEETS:
            for row in self.wb[name].iter_rows():
                for c in row:
                    if isinstance(c.value, str) and c.value.startswith("="):
                        v = self.wbv[name][c.coordinate].value
                        if v is None or v == "":
                            missing.append((name, c.coordinate))
        self.assertEqual(missing, [])

    def test_internal_hyperlinks_resolve(self):
        bad = []
        for ws in self.wb.worksheets:
            for row in ws.iter_rows():
                for c in row:
                    h = c.hyperlink
                    if h is None or not h.location:
                        continue
                    sheet = h.location.split("!")[0].strip("'")
                    if sheet not in self.wb.sheetnames:
                        bad.append((ws.title, c.coordinate, h.location))
        self.assertEqual(bad, [])

    def test_scan_passes(self):
        scan_or_raise(self.data)

    def test_release_values(self):
        self.assertEqual(BUILD.version, "1.0.0")


if __name__ == "__main__":
    unittest.main()
