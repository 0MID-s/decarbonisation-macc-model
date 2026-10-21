"""End-to-end build tests: exit codes, fail-closed behaviour, sheet order."""
import contextlib
import dataclasses
import io
import os
import tempfile
import unittest
from unittest import mock

import openpyxl

import build
from config.schema import Status
from qa.lint import LintError
from qa.readback import ReadbackError
from qa.scan import ScanError, scan_xlsx_bytes


def _quiet_main(argv):
    with contextlib.redirect_stdout(io.StringIO()), \
            contextlib.redirect_stderr(io.StringIO()):
        return build.main(argv)


class BuildTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)

    def path(self, *p):
        return os.path.join(self.tmp.name, *p)

    def test_config_build(self):
        out = self.path("c.xlsx")
        self.assertEqual(_quiet_main(["--dataset", "config", "--out", out]), 0)
        self.assertTrue(os.path.exists(out))

    def test_fixtures_build(self):
        out = self.path("f.xlsx")
        self.assertEqual(_quiet_main(["--dataset", "fixtures", "--out", out]), 0)
        self.assertTrue(os.path.exists(out))

    def test_fixtures_never_written_to_deliverable(self):
        # WR-08: synthetic data must not land on dist/maison_veritas.xlsx.
        cwd = os.getcwd()
        os.chdir(self.tmp.name)
        self.addCleanup(os.chdir, cwd)
        self.assertEqual(_quiet_main(["--dataset", "fixtures"]), 0)
        self.assertTrue(os.path.exists(build.SYNTHETIC_MAIN_OUT))
        self.assertFalse(os.path.exists(build.DELIVERABLE))
        with self.assertRaises(SystemExit) as cm:
            _quiet_main(["--dataset", "fixtures", "--out", build.DELIVERABLE])
        self.assertNotEqual(cm.exception.code, 0)
        self.assertFalse(os.path.exists(build.DELIVERABLE))
        self.assertEqual(_quiet_main(["--target", "spike", "--dataset", "fixtures"]), 0)
        self.assertTrue(os.path.exists(os.path.join("dist", "spike_macc.xlsx")))

    def test_placeholder_blocks_and_writes_nothing(self):
        ds = build.load_dataset("fixtures")
        bad = dataclasses.replace(ds.inputs[3], status=Status.PLACEHOLDER)
        ds = dataclasses.replace(ds, inputs=ds.inputs[:3] + (bad,) + ds.inputs[4:])
        out = self.path("p.xlsx")
        with mock.patch.object(build, "load_dataset", return_value=ds):
            self.assertEqual(_quiet_main(["--out", out]), 2)
        self.assertFalse(os.path.exists(out))

    def test_lint_failure(self):
        def extra(wb, styles, cells):
            ws = wb.add_worksheet("Bad")
            cells.formula(ws, 0, 0, "=A1*1000", None, value=0)
        with self.assertRaises(LintError):
            build.build_bytes("main", build.load_dataset("fixtures"),
                              extra_builders=[extra])

    def test_scan_failure_on_bypass(self):
        def extra(wb, styles, cells):
            wb.add_worksheet("Raw").write_string(0, 0, "=1+1")
        with self.assertRaises(ScanError):
            build.build_bytes("main", build.load_dataset("fixtures"),
                              extra_builders=[extra])

    def test_readback_catches_unregistered_formula(self):
        def extra(wb, styles, cells):
            wb.add_worksheet("Raw").write_formula(0, 0, "=1+1", None, 2)
        with self.assertRaises(ReadbackError) as cm:
            build.build_bytes("main", build.load_dataset("fixtures"),
                              extra_builders=[extra])
        self.assertTrue(any("not in registry" in p for p in cm.exception.problems))

    def test_readback_catches_unregistered_defined_name(self):
        # CR-01: a name written outside CellWriter is a read-back problem.
        def extra(wb, styles, cells):
            wb.define_name("Raw_Name", "=Inputs!$A$1")
        with self.assertRaises(ReadbackError) as cm:
            build.build_bytes("main", build.load_dataset("fixtures"),
                              extra_builders=[extra])
        self.assertTrue(any("not written via CellWriter: Raw_Name" in p
                            for p in cm.exception.problems))

    def test_unreadable_workbook_is_readback_error(self):
        # WR-02: an openpyxl load failure becomes exit code 5, not a traceback.
        default_out, original = build.TARGETS["main"]

        def builder(wb, styles, cells, dataset, cfg):
            layouts = original(wb, styles, cells, dataset, cfg)
            wb.add_worksheet("Raw").write_formula(0, 0, "=A2", None, float("nan"))
            return layouts
        out = self.path("nan.xlsx")
        with mock.patch.dict(build.TARGETS, {"main": (default_out, builder)}):
            with self.assertRaises(ReadbackError) as cm:
                build.build_bytes("main", build.load_dataset("fixtures"))
            self.assertEqual(_quiet_main(["--dataset", "fixtures", "--out", out]), 5)
        self.assertTrue(any("could not read" in p for p in cm.exception.problems))
        self.assertFalse(os.path.exists(out))

    def test_empty_string_cached_value_accepted(self):
        # WR-03: =IF(..,"","") with value="" is a legitimate cached value.
        def extra(wb, styles, cells):
            ws = wb.add_worksheet("Empty")
            ws.hide_gridlines(2)
            ws.freeze_panes(1, 1)
            cells.formula(ws, 2, 2, '=IF(B1,"","")', None, value="")
        build.build_bytes("main", build.load_dataset("fixtures"),
                          extra_builders=[extra])

    def test_main_workbook_clean_and_ordered(self):
        data, _ = build.build_bytes("main", build.load_dataset("config"))
        self.assertEqual(scan_xlsx_bytes(data), [])
        wb = openpyxl.load_workbook(io.BytesIO(data))
        self.assertEqual(wb.sheetnames,
                         ["Cover", "Inputs", "Sources", "Lists", "Baseline", "Rates", "Carbon", "Measures", "MACC", "Notes"])
        self.assertEqual(list(build.main_sheets(build.load_dataset("config"))),
                         wb.sheetnames)

    def test_fixtures_main_has_no_baseline(self):
        data, _ = build.build_bytes("main", build.load_dataset("fixtures"))
        wb = openpyxl.load_workbook(io.BytesIO(data))
        self.assertNotIn("Baseline", wb.sheetnames)

    def test_creates_missing_output_directories(self):
        out = self.path("missing", "nested", "out.xlsx")
        self.assertEqual(_quiet_main(["--dataset", "fixtures", "--out", out]), 0)
        self.assertTrue(os.path.exists(out))


if __name__ == "__main__":
    unittest.main()
