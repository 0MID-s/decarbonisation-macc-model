"""Disclaimer, Cover and Notes shell tests (BASE-06, D-33)."""
import dataclasses
import io
import re
import unittest

import openpyxl

import build
from config.schema import Status
from xl.sheets.common import FICTITIOUS_NOTE, disclaimer_text

EXPECTED = ("‘Maison Veritas’ is a fictitious company. Its baseline "
            "is a representative Belgian NACE C20 mid-sized manufacturer "
            "derived from official Eurostat statistics (sbs_sc_ovw, "
            "nrg_bal_c, env_ac_ainah_r2, retrieved %s). It does not "
            "represent any real company.")


class Notes(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.ds = build.load_dataset("config")
        data, _ = build.build_bytes("main", cls.ds)
        cls.wb = openpyxl.load_workbook(io.BytesIO(data))

    def test_disclaimer_exact(self):
        latest = max(s.retrieved for s in self.ds.sources
                     if s.id in ("Src_EurostatSBS", "Src_EurostatEnergyBalance",
                                 "Src_EurostatAirAccounts"))
        self.assertEqual(disclaimer_text(self.ds.sources),
                         EXPECTED % latest.isoformat())

    def test_fallback_for_missing_sources(self):
        fx = build.load_dataset("fixtures")
        self.assertEqual(disclaimer_text(fx.sources), FICTITIOUS_NOTE)

    def test_cover_and_notes_same_string(self):
        want = disclaimer_text(self.ds.sources)
        self.assertEqual(self.wb["Cover"]["B5"].value, want)
        self.assertEqual(self.wb["Notes"]["B5"].value, want)

    def test_headings(self):
        vals = [c.value for row in self.wb["Notes"].iter_rows() for c in row]
        for h in ("About this model", "Data scope and conventions",
                  "Limitations", "Exclusions", "Carbon price notes",
                  "Transcribed inputs (D-14 route 3)", "Version log"):
            self.assertIn(h, vals)

    def test_phase3_section(self):
        vals = [c.value for row in self.wb["Notes"].iter_rows() for c in row]
        self.assertIn("Measures and MACC method (Phase 3)", vals)
        t = self._text()
        for w in ("not an optimiser", "Kesicki", "stack", "shadow-price",
                  "Phase 3: measures engine"):
            self.assertIn(w, t)
        fx = build.load_dataset("fixtures")
        data, _ = build.build_bytes("main", fx)
        import io, openpyxl
        fvals = [c.value for row in openpyxl.load_workbook(
            io.BytesIO(data))["Notes"].iter_rows() for c in row]
        self.assertNotIn("Measures and MACC method (Phase 3)", fvals)

    def test_phase4_section(self):
        vals = [c.value for row in self.wb["Notes"].iter_rows() for c in row]
        self.assertIn("Carbon exposure and sensitivity method (Phase 4)", vals)
        t = " ".join(self._text().split())
        for w in ("not a live cell", "Python-verified at build",
                  "upper end of the literature", "Phase 4: carbon cost exposure"):
            self.assertIn(w, t)
        fx = build.load_dataset("fixtures")
        data, _ = build.build_bytes("main", fx)
        fvals = [c.value for row in openpyxl.load_workbook(
            io.BytesIO(data))["Notes"].iter_rows() for c in row]
        self.assertNotIn("Carbon exposure and sensitivity method (Phase 4)",
                         fvals)

    def _text(self):
        return "\n".join(str(c.value) for row in self.wb["Notes"].iter_rows()
                         for c in row if c.value is not None)

    def test_transcribed_list_generated(self):
        t = self._text()
        for r in self.ds.inputs:
            if (r.note.startswith("Transcribed")
                    or r.rationale.startswith(("Transcribed",
                                           "Unverified transcription"))):
                self.assertIn(r.id + ":", t)
        self.assertIn("Carbon_EUA_Avg2025", t)

    def test_no_projection_wording(self):
        self.assertIn("Not captured: no official projection located on "
                      "allowlisted sources (see RESEARCH 3.11)", self._text())

    def test_projection_captured_wording(self):
        from xl.sheets.notes import ets2_projection_line
        rec = dataclasses.replace(
            [r for r in self.ds.inputs if r.id == "Carbon_EUA_Avg2025"][0],
            id="Carbon_ETS2_Proj_2030")
        line = ets2_projection_line(self.ds.inputs + (rec,), self.ds.sources)
        self.assertIn("captured as data:", line)
        self.assertIn("not used in the model calculations", line)
        self.assertIn("Carbon_ETS2_Proj_2030", line)

    def test_flat_grid_deviation_stated(self):
        self.assertIn("DATA-05", self._text())
        self.assertIn("D-23", self._text())

    def test_no_formula_leading_chars(self):
        for row in self.wb["Notes"].iter_rows():
            for c in row:
                if isinstance(c.value, str):
                    self.assertNotIn(c.value[:1], "=+-@")

    def test_prose_has_no_input_values(self):
        import xl.sheets.notes as nt
        from checkmodel.baseline_reference import model_values
        vals = set(model_values(self.ds.inputs).values())
        prose = " ".join(nt.ABOUT + nt.CONVENTIONS + nt.LIMITATIONS
                         + nt.EXCLUSIONS + nt.METHOD_CHOICES
                         + nt.LIMITATIONS_P5
                         + (nt.legal_status_note(self.ds.inputs),))
        prose = re.sub(r"(Phase|Scope) \d", "", prose)
        # bibliographic citation, not a model figure
        prose = re.sub(r"Climate Policy 12\(2\), 219-236, doi:\S+", "", prose)
        for tok in re.findall(r"(?<![\w.-])\d+(?:\.\d+)?(?![\w.-])", prose):
            x = float(tok)
            if 1900 <= x <= 2100:
                continue
            self.assertNotIn(x, vals, tok)

    def test_legal_status_note_follows_config(self):
        """IN-05: the Notes legal paragraph is generated from the statuses."""
        import xl.sheets.notes as nt
        from config.schema import Status
        text = nt.legal_status_note(self.ds.inputs)
        self.assertIn(" ".join(text.split()), " ".join(self._text().split()))
        legal = [r for r in self.ds.inputs
                 if r.id.startswith(nt.LEGAL_PREFIXES)
                 or r.id in nt.LEGAL_EXTRA_IDS]
        self.assertGreaterEqual(len(legal), 16)
        head, _, tail = text.partition("Analyst assumptions")
        for r in legal:
            self.assertIn(r.id, head if r.status is Status.SOURCED else tail,
                          r.id)
        self.assertIn("Dir. (EU) 2023/959 as enacted", text)
        self.assertIn("original 2003 Annex I wording", text)

    def test_phase5_sections(self):
        vals = [c.value for row in self.wb["Notes"].iter_rows() for c in row]
        for h in ("Method choices", "Limitations (Phase 5)",
                  "Disclosure and Dashboard corrections",
                  "Excel verification record"):
            self.assertIn(h, vals)
        t = " ".join(self._text().split())
        for w in ("10.1080/14693062.2011.582347", "pre-tax", "ex-VAT",
                  "independent", "cumulative and undiscounted"):
            self.assertIn(w, t)
        self.assertNotIn("full method and version log in Phase 5", t)

    def test_no_stale_phase_language(self):
        # WR-09: the v1.0 Notes describe what was built, not pending work
        import xl.sheets.notes as nt
        t = " ".join(self._text().split())
        for bad in ("to be resolved", "candidate Phase", "Phase 4 decides",
                    "belong to Phase", "to test in Phase",
                    "captured as data for Phase", "not used in Phase"):
            self.assertNotIn(bad, t)
        line = " ".join(nt.ets2_projection_line(
            self.ds.inputs, self.ds.sources).split())
        self.assertNotIn("Phase", line)
        self.assertIn("excluded from capex and shown as a memo line", t)
        self.assertIn("tornado on the Sensitivity sheet", t)

    # WR-14: phase-planning phrases that describe pending work rather than
    # the built state. Section headings such as "(Phase 3)", the phase
    # history list and "DATA_LOG Phase 3" document references are allowed.
    PLANNING = re.compile(
        r"Phase \d+ (?:sensitivity|decision|decides|tests|inflation|X_VAT|D-)"
        r"|(?:candidate|in|for|to test in|belong to|not used in) Phase \d"
        r"|resolved against|to be resolved|captured as data for Phase")

    def test_inputs_rationale_has_no_planning_language(self):
        # WR-14: every Inputs rationale and note describes the built state.
        for r in self.ds.inputs:
            for field in ("rationale", "note"):
                text = " ".join((getattr(r, field) or "").split())
                self.assertIsNone(self.PLANNING.search(text), (r.id, field))
                bare = re.sub(r"(?:DATA_LOG|NGFS) Phase \d", "", text)
                self.assertNotRegex(bare, r"Phase \d", (r.id, field))
        # and the rendered Rationale / notes cells on the Inputs sheet
        ws = self.wb["Inputs"]
        hdr = [c.value for c in ws[11]]
        self.assertIn("Rationale / notes", hdr)
        col = hdr.index("Rationale / notes") + 1
        n = 0
        for row in ws.iter_rows(min_row=12, min_col=col, max_col=col):
            v = row[0].value
            if isinstance(v, str):
                n += 1
                self.assertIsNone(self.PLANNING.search(" ".join(v.split())),
                                  row[0].coordinate)
        self.assertGreater(n, 100)

    def test_client_sheets_have_no_planning_language(self):
        # WR-14: Notes, Rates and Carbon text cells, the same rule
        for name in ("Notes", "Rates", "Carbon"):
            for row in self.wb[name].iter_rows():
                for c in row:
                    if isinstance(c.value, str) and not c.value.startswith("="):
                        self.assertIsNone(
                            self.PLANNING.search(" ".join(c.value.split())),
                            "%s!%s" % (name, c.coordinate))

    def test_inputs_sensitivity_claims_are_true(self):
        # WR-14: only inputs that the Sensitivity sheet really varies say so
        by = {r.id: r for r in self.ds.inputs}
        not_varied = ("Meas_LED_LightingShare", "Meas_LED_Coverage",
                      "Meas_LED_SavingRate", "Meas_VSD_MotorShare",
                      "Meas_VSD_Coverage", "Meas_VSD_SavingRate",
                      "Meas_PV_Size", "Meas_Fleet_MaintSaving",
                      "Meas_Fleet_DieselVanPrice", "Meas_RC_Term")
        for k in not_varied:
            self.assertEqual(by[k].note, "Analyst assumption; not varied on "
                             "the Sensitivity sheet", k)
        self.assertIn("Capex tornado driver", by["Meas_Fleet_ChargerCapex"].note)
        self.assertNotIn("6,000 h", by["Meas_HP_FullLoadHours"].rationale)
        for k in ("Meas_HP_IndirectCapex_Trad", "Meas_HP_IndirectCapex_HT"):
            self.assertIn("excluded from capex and shown as a memo line",
                          by[k].rationale)

    def test_measures_assumptions_stated(self):
        # WR-07: full-tariff electricity valuation and the CO2-only residual
        # mix are stated in Notes
        t = " ".join(self._text().split())
        for w in ("full Eurostat X_VAT band price (BL_ElecPrice)",
                  "the PV saving is an upper bound",
                  "(EF_ResidualMB, AIB) is CO2 only"):
            self.assertIn(w, t)

    def test_corrections_rendered(self):
        from config.disclosure import CORRECTIONS
        vals = [c.value for row in self.wb["Notes"].iter_rows() for c in row]
        for row in CORRECTIONS:
            for cell in row:
                self.assertIn(cell, vals)

    def test_no_citation_hyperlinks(self):
        for row in self.wb["Notes"].iter_rows():
            for c in row:
                if c.hyperlink is not None:
                    tgt = str(c.hyperlink.target or "") + str(c.hyperlink.location or "")
                    self.assertNotIn("tandfonline.com", tgt)
                    self.assertNotIn("doi.org", tgt)

    def test_version_log_table(self):
        from config.build_config import VERSION_LOG
        ws = self.wb["Notes"]
        rows = [tuple(ws.cell(r, c).value for c in (2, 3, 4))
                for r in range(1, ws.max_row + 1)]
        i = rows.index(("Version", "Date", "Change"))
        got = rows[i + 1:i + 1 + len(VERSION_LOG)]
        self.assertEqual(tuple(got), tuple(VERSION_LOG))

    def test_excel_record_honest(self):
        from config.build_config import EXCEL_RECORD
        t = " ".join(self._text().split())
        for line in EXCEL_RECORD:
            self.assertIn(line, t)
        self.assertIn("Excel for the web", t)
        if any("not verified" in l for l in EXCEL_RECORD):
            self.assertNotIn("verified in Excel for the web", t)


if __name__ == "__main__":
    unittest.main()
