"""MACC sheet tests (plan 03-09): live sort, tiers, reference price, charts,
checked against the reference model in all six selector states."""
import io
import re
import unittest
import zipfile

import openpyxl

import build
from checkmodel import appraisal as ap
from checkmodel.measures_reference import (
    BASES, DEFAULT_STATE, SCENARIOS_UI, compute_measures, macc_table)
from config.inputs import INPUTS
from config.measures_catalog import MEASURES
from qa.formula_eval import evaluate_names

NAMES = ["MACC_RefPrice", "MACC_Title", "MACC_Rank_Measure", "MACC_Rank_Cost",
         "MACC_Rank_Width", "MACC_Rank_Tier", "MACC_Chart_Width",
         "MACC_Chart_Height"]
TIER_LABEL = {"Do first": "Do first", "Next": "Next", "Later": "Later"}


def _close(a, b):
    return abs(a - b) <= 1e-9 * max(1.0, abs(b))


class MaccSheet(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        ds = build.load_dataset("config")
        cls.data, cls.registry = build.build_bytes("main", ds)
        cls.wb = openpyxl.load_workbook(io.BytesIO(cls.data))
        cls.wbv = openpyxl.load_workbook(io.BytesIO(cls.data), data_only=True)

    def _check_state(self, out, basis, scen):
        ref = compute_measures(INPUTS, basis, scen)
        table = macc_table(ref)
        self.assertTrue(_close(out["MACC_RefPrice"], ref.ref_price))
        self.assertEqual(out["MACC_Title"], "MACC: %s basis, %s carbon price "
                         "scenario" % (basis, scen))
        self.assertEqual(out["MACC_Rank_Measure"], [t[0] for t in table])
        self.assertEqual(out["MACC_Rank_Tier"], [t[4] for t in table])
        for got, (_n, w, c, elig, _t) in zip(out["MACC_Rank_Cost"], table):
            if elig:
                self.assertTrue(_close(got, c), (got, c))
            else:
                self.assertEqual(got, ap.NA_NO_ABATEMENT)
        for got, (_n, w, c, elig, _t) in zip(out["MACC_Rank_Width"], table):
            self.assertTrue(_close(got, w), (got, w))
        sorted_idx = [[m.key for m in ref.measures].index(k)
                      for k in ref.sorted_keys]
        self.assertEqual(out["MACC_Chart_Width"],
                         [ref.chart_widths[i] for i in sorted_idx])
        for got, (_n, w, c, elig, _t) in zip(out["MACC_Chart_Height"], table):
            self.assertTrue(_close(got, c if elig else 0), (got, c))
        return ref

    def test_default_state_cached(self):
        ref = compute_measures(INPUTS, *DEFAULT_STATE)
        table = macc_table(ref)
        out = {}
        for n in NAMES:
            sheet, a = self.wb.defined_names[n].attr_text.split("!")
            cells = self.wbv[sheet][a.replace("$", "")]
            if ":" in a:
                out[n] = [c[0].value for c in self.wbv[sheet][a.replace("$", "")]]
            else:
                out[n] = cells.value
        self.assertEqual(out["MACC_Rank_Measure"], [t[0] for t in table])
        self.assertEqual(out["MACC_Rank_Tier"], [t[4] for t in table])
        self.assertTrue(_close(out["MACC_RefPrice"], ref.ref_price))
        costs = [t[2] for t in table if t[3]]
        self.assertEqual(costs, sorted(costs))

    def test_six_states(self):
        for basis in BASES:
            for scen in SCENARIOS_UI:
                with self.subTest(state=(basis, scen)):
                    out = evaluate_names(self.data, NAMES, overrides={
                        ("Inputs", "G8"): scen, ("Inputs", "G9"): basis})
                    ref = self._check_state(out, basis, scen)
                    if basis == "Location-based":
                        table = macc_table(ref)
                        self.assertEqual(table[-1][0], MEASURES[-1].name)
                        self.assertEqual(out["MACC_Rank_Cost"][-1],
                                         ap.NA_NO_ABATEMENT)
                        self.assertEqual(out["MACC_Chart_Width"][-1], 0)
                        self.assertEqual(out["MACC_Rank_Tier"][-1], "Later")

    def test_tiers_move_with_scenario_not_order(self):
        lo = compute_measures(INPUTS, "Market-based", "Low")
        hi = compute_measures(INPUTS, "Market-based", "High")
        self.assertEqual(lo.sorted_keys, hi.sorted_keys)
        ov = lambda s: evaluate_names(self.data, NAMES, overrides={  # noqa: E731
            ("Inputs", "G8"): s, ("Inputs", "G9"): "Market-based"})
        a, b = ov("Low"), ov("High")
        self.assertEqual(a["MACC_Rank_Measure"], b["MACC_Rank_Measure"])
        if lo.tiers != hi.tiers:
            self.assertNotEqual(a["MACC_Rank_Tier"], b["MACC_Rank_Tier"])
        else:
            self.assertEqual(a["MACC_Rank_Tier"],
                             [macc_table(lo)[i][4] for i in range(6)])
        self.assertNotEqual(a["MACC_RefPrice"], b["MACC_RefPrice"])

    def test_chart_xml(self):
        with zipfile.ZipFile(io.BytesIO(self.data)) as z:
            charts = {n: z.read(n).decode("utf-8") for n in z.namelist()
                      if n.startswith("xl/charts/chart")}
        macc = [x for x in charts.values() if "MACC!" in x or "'MACC'!" in x]
        self.assertEqual(len(macc), 2, len(charts))
        a = [x for x in macc if "<c:dateAx>" in x][0]
        self.assertIn("<c:majorUnit", a)
        self.assertRegex(a, r"<c:title>.*MACC!\$B\$5")
        self.assertIn("<c:lineChart>", a)
        self.assertIn("prstDash", a)
        c = [x for x in macc if "<c:dateAx>" not in x][0]
        self.assertIn("<c:barChart>", c)
        self.assertIn("Fallback view", c)

    def test_barrier_notes_labelled(self):
        ws = self.wb["MACC"]
        n = 0
        for row in ws.iter_rows():
            for c in row:
                if c.value in {m.barrier_note for m in MEASURES}:
                    self.assertTrue(c.value.startswith("Analyst commentary:"))
                    n += 1
        self.assertEqual(n, len(MEASURES))
        for r in range(9, 15):
            self.assertTrue(
                self.wbv["MACC"].cell(row=r, column=11).value.startswith(
                    "Analyst commentary:"))

    def test_formulas_literal_free_and_live_sorted(self):
        ws = self.wb["MACC"]
        for r in range(9, 15):
            for col in (3, 4, 5, 6, 7, 8, 9, 11):
                f = ws.cell(row=r, column=col).value
                self.assertTrue(f.startswith("=INDEX("), f)
        lits = re.compile(r"(?<![A-Za-z$_\d.])(\d+\.?\d*)(?![\d:A-Za-z!(])")
        for rec in self.registry:
            if rec.sheet != "MACC":
                continue
            body = re.sub(r'"[^"]*"', '""', rec.text)
            body = re.sub(r"\$?[A-Z]{1,3}\$?\d+", "R", body)
            for m in lits.findall(body):
                self.assertIn(m, ("0", "1"), rec.text)


if __name__ == "__main__":
    unittest.main()
