"""Checks sheet tests (plan 05-03): inventory, deliberate break, selector, gate."""
import io
import re
import unittest
from unittest import mock

import openpyxl

import build
from qa.formula_eval import compare_cached, evaluate_names

AGG = {"M_Check_All", "Expo_Check_All", "Sens_Check_All"}
CHECK_RE = re.compile(r"^(M|Carb|BL|Expo|Sens)_Check_")
ALLOWED = re.compile(
    r"^=(IFERROR\([A-Za-z0-9_]+,FALSE\)|IFERROR\([A-Za-z0-9_]+=\"OK\",FALSE\)"
    r"|IFERROR\(ISNUMBER\(MATCH\([A-Za-z_]+,[A-Za-z_]+,0\)\),FALSE\)"
    r"|IFERROR\(AND\(ABS\(SUMPRODUCT\(MACC_Data_Abatement,MACC_Data_Eligible\)"
    r"-MACC_TotalAbatement\)<=COUNT\(MACC_Chart_Width\)\*Const_MACCRoundTol,"
    r"ABS\(SUM\(MACC_Data_Abatement\)-Dash_AbatedSel\)<=Const_CheckTolerance\),FALSE\)"
    r"|IFERROR\(AND\(Dash_Share>=0,Dash_Share<=1\),FALSE\)"
    r"|IFERROR\(OR\(Dash_AbatedSel=M_Total_tLB,Dash_AbatedSel=M_Total_tMB\),FALSE\)"
    r"|AND\(.*\)|COUNTIF\(.*,FALSE\)|SUM\(.*\)|IF\(AND\(.*\),\"OK\",\"ERROR\"\))$")


def _as_bool(got):
    """Owner text checks read "OK"/"ERROR"; the Checks row tests ="OK"."""
    return dict((k, (v == "OK") if v in ("OK", "ERROR") else v)
                for k, v in got.items())


class ChecksSheet(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.ds = build.load_dataset("config")
        cls.data, _ = build.build_bytes("main", cls.ds)
        cls.wb = openpyxl.load_workbook(io.BytesIO(cls.data))
        cls.ws = cls.wb["Checks"]

    def _addr(self, name):
        sheet, addr = self.wb.defined_names[name].attr_text.split("!")
        return sheet, addr.replace("$", "")

    def _status_formulas(self):
        out = []
        for row in self.ws.iter_rows(min_col=4, max_col=4):
            c = row[0]
            if isinstance(c.value, str) and c.value.startswith("=IFERROR("):
                out.append(c.value)
        return out

    def test_sheet_before_notes(self):
        names = self.wb.sheetnames
        self.assertLess(names.index("Checks"), names.index("Notes"))

    def test_inventory_complete_once(self):
        forms = self._status_formulas()
        expected = [n for n in self.wb.defined_names
                    if CHECK_RE.match(n) and n not in AGG]
        self.assertGreaterEqual(len(expected), 56)
        for n in expected:
            hits = [f for f in forms
                    if re.search(r"IFERROR\(%s[,=]" % re.escape(n), f)]
            self.assertEqual(len(hits), 1, n)
        for f in forms:
            self.assertNotIn("_Check_All", f)

    def test_formula_shapes_only(self):
        for row in self.ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    self.assertRegex(c.value, ALLOWED, c.coordinate)

    def test_default_ok_and_cached(self):
        got = evaluate_names(self.data, ["Chk_Master", "Chk_FailCount"])
        self.assertEqual(got, {"Chk_Master": "OK", "Chk_FailCount": 0})
        bad, _ = compare_cached(self.data)
        self.assertEqual(bad, [])

    def test_deliberate_break(self):
        sheet, addr = self._addr("Meas_PV_Size")
        got = evaluate_names(self.data, ["Chk_Master", "Chk_FailCount"],
                             overrides={(sheet, addr): 100000.0})
        self.assertEqual(got["Chk_Master"], "ERROR")
        self.assertGreaterEqual(got["Chk_FailCount"], 1)

    def test_invalid_selector(self):
        sheet, addr = self._addr("Sel_Basis")
        got = evaluate_names(self.data, ["Chk_Master", "Chk_Sel_Basis"],
                             overrides={(sheet, addr): "Foo"},
                             tolerate_errors=True)
        self.assertEqual(got["Chk_Master"], "ERROR")
        self.assertIs(got["Chk_Sel_Basis"], False)

    def test_dashboard_checks(self):
        names = ["Chk_TileVsMACC", "Chk_ShareInRange", "Chk_BasisExact"]
        for n in names:
            self.assertIn(n, self.wb.defined_names)
        got = evaluate_names(self.data, names)
        self.assertEqual(got, {n: True for n in names})
        # the group 6 Boolean row sits before the text row
        sheet, addr = self._addr("Chk_BasisExact")
        row = int(re.sub(r"\D", "", addr))
        self.assertEqual(self.ws["D%d" % (row + 1)].value, "Enforced at build")

    def test_tile_vs_macc_fails_at_zero_tolerance(self):
        sheet, addr = self._addr("Const_MACCRoundTol")
        got = evaluate_names(self.data, ["Chk_TileVsMACC"],
                             overrides={(sheet, addr): 0.0})
        self.assertIs(got["Chk_TileVsMACC"], False)

    def test_negative_abatement_is_not_a_master_error(self):
        # CR-01: a measure with negative abatement is a valid state. The MACC
        # leaves it out, the tile keeps it (net sum), and the master stays OK.
        sheet, addr = self._addr("Meas_Fleet_EVkWhPerKm")
        names = ["Chk_Master", "Chk_FailCount", "Chk_TileVsMACC",
                 "M_Check_All", "Expo_Check_All", "Sens_Check_All",
                 "M_Fleet_AvgtSel", "M_Fleet_Eligible", "Dash_AbatedSel",
                 "MACC_TotalAbatement"]
        got = evaluate_names(self.data, names,
                             overrides={(sheet, addr): 1.5})
        self.assertLess(got["M_Fleet_AvgtSel"], 0)
        self.assertEqual(got["M_Fleet_Eligible"], 0)
        # the two totals really differ, so the old check would have failed
        self.assertGreater(abs(got["Dash_AbatedSel"]
                               - got["MACC_TotalAbatement"]), 6 * 0.5)
        self.assertIs(got["Chk_TileVsMACC"], True)
        for n in ("M_Check_All", "Expo_Check_All", "Sens_Check_All"):
            self.assertIs(got[n], True, n)
        self.assertEqual(got["Chk_Master"], "OK")
        self.assertEqual(got["Chk_FailCount"], 0)

    def test_rising_grid_factor_gives_no_negative_width(self):
        # WR-03: EF_GridLB_Change at +8%/yr and +10%/yr under location-based
        # once gave the heat pump / fleet a negative MACC width. One rule
        # (PV of tonnes and average tonnes both > 0) now drives the flag.
        from checkmodel.measures_reference import compute_measures
        g = self._addr("EF_GridLB_Change")
        b = self._addr("Sel_Basis")
        from checkmodel.outputs_reference import compute_outputs
        names = ["Chk_Master", "Chk_FailCount", "MACC_Chart_Width",
                 "M_HP_Eligible", "M_Fleet_Eligible", "MACC_TotalAbatement"]
        for x, key in ((0.08, "HP"), (0.10, "Fleet")):
            ov = {"EF_GridLB_Change": x}
            out = compute_outputs(self.ds.inputs, "Location-based", "Central",
                                  "Off", overrides=ov)
            got = evaluate_names(self.data, names + list(out.checks),
                                 overrides={g: x, b: "Location-based"})
            widths = got["MACC_Chart_Width"]
            self.assertTrue(all(w >= 0 for w in widths), (x, widths))
            self.assertEqual(got["M_%s_Eligible" % key], 0)
            ref = compute_measures(self.ds.inputs, "Location-based", "Central",
                                   overrides=ov)
            self.assertEqual(sorted(widths), sorted(ref.chart_widths))
            self.assertEqual(got["MACC_TotalAbatement"], sum(ref.chart_widths))
            # every MACC / tile / eligibility check passes. The only failures
            # are the annuity reconciliations, which hold only for flat tonne
            # profiles (a rising grid factor is not flat); the reference
            # model agrees name by name, so Chk_Master is ERROR in both.
            failing = sorted(n for n, v in _as_bool(
                {n: got[n] for n in out.checks}).items() if v is not True)
            self.assertTrue(failing)
            self.assertTrue(all(n.startswith("M_Check_")
                                and n.endswith("_Annuity") for n in failing),
                            failing)
            self.assertEqual(failing, sorted(n for n, v in out.checks.items()
                                             if not v))
            self.assertIs(got["Chk_TileVsMACC"], True)
            self.assertEqual(got["Chk_Master"], out.chk_master)
            self.assertEqual(got["Chk_FailCount"], out.chk_fail_count)

    def test_break_even_check_can_fail(self):
        # WR-08: the break-even check recomputes NPV incl. carbon value at
        # the break-even price from the yearly rows. A wrong NPV cell (which
        # both the levelised cost and the break-even formula read) used to
        # pass the old "break-even = levelised" self-comparison; now it fails.
        sheet, addr = self._addr("M_HP_NPV")
        names = ["M_Check_HP_BreakEven", "M_Check_LED_BreakEven",
                 "M_HP_LevCost", "M_HP_BreakEven", "Chk_Master"]
        base = evaluate_names(self.data, names)
        self.assertIs(base["M_Check_HP_BreakEven"], True)
        got = evaluate_names(self.data, names, overrides={(sheet, addr): 0.0})
        self.assertEqual(got["M_HP_LevCost"], got["M_HP_BreakEven"])
        self.assertIs(got["M_Check_HP_BreakEven"], False)
        self.assertIs(got["M_Check_LED_BreakEven"], True)
        self.assertEqual(got["Chk_Master"], "ERROR")
        f = self.wb["Checks"]
        desc = [c.value for row in f.iter_rows() for c in row
                if isinstance(c.value, str) and "break-even price" in c.value]
        self.assertTrue(desc)
        self.assertFalse(any("equals levelised cost" in d for d in desc))

    def test_tile_vs_macc_catches_tile_drift(self):
        # the second leg ties the tile to the per-measure MACC data
        sheet, addr = self._addr("Dash_AbatedSel")
        got = evaluate_names(self.data, ["Chk_TileVsMACC"],
                             overrides={(sheet, addr): 1.0})
        self.assertIs(got["Chk_TileVsMACC"], False)

    def test_cached_values_from_reference(self):
        # WR-01: each status row caches the reference value, not a literal,
        # and the inventory matches the reference model one to one.
        from checkmodel.outputs_reference import compute_outputs
        ref = compute_outputs(self.ds.inputs, "Market-based", "Central", "Off")
        names = [n for n in ref.checks]
        rows = [n for n in self.wb.defined_names
                if CHECK_RE.match(n) and n not in AGG]
        rows += [n for n in self.wb.defined_names if n.startswith("Chk_")
                 and n not in ("Chk_Master", "Chk_FailCount")]
        self.assertEqual(sorted(names), sorted(rows))
        got = evaluate_names(self.data, names)
        self.assertEqual(_as_bool(got), ref.checks)
        # and the Checks status cells cache exactly those values
        wbv = openpyxl.load_workbook(io.BytesIO(self.data), data_only=True)
        cached = [c.value for row in wbv["Checks"].iter_rows(min_col=4,
                                                             max_col=4)
                  for c in row if isinstance(c.value, bool)]
        self.assertEqual(len(cached), len(names))
        self.assertEqual(cached, [True] * len(names))

    def test_reference_matches_formulas_in_failing_state(self):
        # WR-01: in a state where checks legitimately fail, the reference
        # inventory and the workbook formulas agree name by name.
        from checkmodel.outputs_reference import compute_outputs
        ref = compute_outputs(self.ds.inputs, "Market-based", "Central", "Off",
                              overrides={"Meas_PV_Size": 100000.0})
        sheet, addr = self._addr("Meas_PV_Size")
        names = list(ref.checks) + ["Chk_Master", "Chk_FailCount"]
        got = evaluate_names(self.data, names,
                             overrides={(sheet, addr): 100000.0})
        self.assertEqual(got["Chk_Master"], ref.chk_master)
        self.assertEqual(got["Chk_Master"], "ERROR")
        self.assertEqual(got["Chk_FailCount"], ref.chk_fail_count)
        got = _as_bool(got)
        for n in ref.checks:
            self.assertEqual(got[n], ref.checks[n], n)

    def test_gate_compares_cached_with_evaluated(self):
        # WR-01: the cached half of the gate can now fail.
        with mock.patch.object(build, "evaluate_names",
                               return_value={"Chk_Master": "OK",
                                             "Chk_FailCount": 3}):
            with self.assertRaises(build.BuildGateError) as cm:
                build.build_bytes("main", self.ds)
        self.assertIn("Chk_FailCount", str(cm.exception))

    def test_group6_text_row(self):
        hit = [c for row in self.ws.iter_rows(min_col=4, max_col=4)
               for c in row if c.value == "Enforced at build"]
        self.assertEqual(len(hit), 1)
        self.assertEqual(hit[0].data_type, "s")

    def test_gate(self):
        with mock.patch.object(build, "evaluate_names",
                               return_value={"Chk_Master": "ERROR"}):
            with self.assertRaises(build.BuildGateError):
                build.build_bytes("main", self.ds)
            with mock.patch("sys.stderr"):
                self.assertEqual(
                    build.main(["--out", "/dev/null/x.xlsx"]), 6)


if __name__ == "__main__":
    unittest.main()
