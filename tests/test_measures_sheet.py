"""Measures sheet tests (plan 03-07): formulas, cached values vs the
reference model, tags, checks, and selector behaviour."""
import io
import unittest

import openpyxl

import build
from checkmodel.measures_reference import compute_measures
from config.inputs import INPUTS
from config.measures_catalog import MEASURES
from qa.formula_eval import compare_cached, evaluate_names
from qa.lint import lint_formula
from xl.layout import TIMELINE_FIRST_COL, TIMELINE_LAST_COL

KEYS = [m.key for m in MEASURES]
REL = 1e-9


class MeasuresSheet(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        ds = build.load_dataset("config")
        cls.data, cls.registry = build.build_bytes("main", ds)
        cls.wb = openpyxl.load_workbook(io.BytesIO(cls.data))
        cls.wbv = openpyxl.load_workbook(io.BytesIO(cls.data), data_only=True)
        cls.ref = compute_measures(INPUTS, "Market-based", "Central")

    def _cells(self, wb, name):
        sheet, addr = self.wb.defined_names[name].attr_text.split("!")
        a = addr.replace("$", "")
        ws = wb[sheet]
        if ":" not in a:
            return [ws[a]]
        first = a.split(":")[0]
        r = ws[first].row
        return [ws.cell(row=r, column=c + 1)
                for c in range(TIMELINE_FIRST_COL, TIMELINE_LAST_COL + 1)]

    def _vals(self, name):
        return [c.value for c in self._cells(self.wbv, name)]

    def _first_formula(self, name):
        return self._cells(self.wb, name)[0].value

    def _assert_rows(self, name, want):
        got = self._vals(name)
        self.assertEqual(len(got), len(want), name)
        for g, w in zip(got, want):
            self.assertLessEqual(abs(g - w), REL * max(1.0, abs(w)), name)

    def test_formula_strings(self):
        idx_row = self._cells(self.wb, "M_Idx")[0].row
        self.assertEqual(self._first_formula("M_DF"),
                         "=1/(1+Rate_DiscountReal)^N$%d" % idx_row)
        sel = self._first_formula("M_EF_Sel")
        self.assertTrue(sel.startswith("=CHOOSE(M_BasisIdx,"), sel)
        self.assertEqual(self._first_formula("M_BasisIdx"),
                         "=MATCH(Sel_Basis,List_Basis,0)")
        self.assertEqual(self._first_formula("M_ScenIdx"),
                         "=MATCH(Sel_Scenario,List_Scenario,0)")
        cov_row = self._cells(self.wb, "M_RC_Covered")[0].row
        tlb = self._first_formula("M_RC_tLB")
        self.assertNotIn("N$%d" % cov_row, tlb)
        self.assertIn("N$%d" % cov_row, self._first_formula("M_RC_tMB"))
        capex_row = self._cells(self.wb, "M_LED_CF")[0].row - 7
        led_capex = self.wb["Measures"].cell(
            row=capex_row, column=TIMELINE_FIRST_COL + 1).value
        self.assertTrue(led_capex.startswith("=-$D$"), led_capex)
        self.assertIn("Meas_LED_ZMCapex_EUR2025",
                      self._cells(self.wb, "M_LED_Capex")[0].value)

    def test_cached_values_equal_reference(self):
        ref = self.ref
        self._assert_rows("M_DF", ref.df)
        self._assert_rows("M_PriceIdx", ref.price_idx)
        self._assert_rows("M_CarbonPrice", ref.carbon)
        self._assert_rows("M_EF_LB", ref.f_lb)
        self._assert_rows("M_EF_MB", ref.f_mb)
        self._assert_rows("M_EF_Sel", ref.f_sel)
        self._assert_rows("M_Op", ref.op)
        self._assert_rows("M_Idx", list(range(len(ref.years))))
        st = ref.stack
        for n, want in (("S0", st.s0), ("S1", st.s1), ("S2", st.s2),
                        ("S3", st.s3), ("S4", st.s4), ("S5", st.s5),
                        ("GasAfter", st.gas_after),
                        ("DieselAfter", st.diesel_after)):
            self._assert_rows("M_Stack_" + n, want)
        for m in ref.measures:
            k = m.key
            for suffix, want in (
                    ("dElec", m.d_elec), ("dGas", m.d_gas),
                    ("dDiesel", m.d_diesel), ("Covered", m.covered),
                    ("tLB", m.t_lb), ("tMB", m.t_mb), ("tSel", m.t_sel),
                    ("CF", m.cf_fin), ("CFCarbon", m.cf_incl)):
                self._assert_rows("M_%s_%s" % (k, suffix), want)
            self.assertAlmostEqual(self._vals("M_%s_Capex" % k)[0], m.capex,
                                   delta=REL * max(1.0, abs(m.capex)))
            self.assertEqual(self._vals("M_%s_Life" % k)[0], m.life)
            self.assertAlmostEqual(self._vals("M_%s_ResFrac" % k)[0],
                                   m.res_frac, delta=1e-12)

    def test_checks_true(self):
        names = [n for n in self.wb.defined_names if n.startswith("M_Check_")]
        self.assertGreaterEqual(len(names), 7)
        for n in ("M_Check_StackNonNeg", "M_Check_ContractCap",
                  "M_Check_PVSelfConsumed", "M_Check_ElecReductionsCap",
                  "M_Check_GasCap", "M_Check_DieselCap",
                  "M_Check_PVPeakBelowLoad"):
            self.assertIn(n, names)
        for n in names:
            self.assertIs(self._vals(n)[0], True, n)

    def test_names_for_phase4(self):
        for k in KEYS:
            for s in ("dElec", "dGas", "dDiesel", "Covered", "tLB", "tMB",
                      "tSel", "CF", "CFCarbon"):
                self.assertIn("M_%s_%s" % (k, s), self.wb.defined_names)
        for n in ("M_Idx", "M_Op", "M_DF", "M_PriceIdx", "M_CarbonPrice",
                  "M_EF_LB", "M_EF_MB", "M_EF_Sel", "M_LastIdx",
                  "M_BasisIdx", "M_ScenIdx", "M_DieselExVAT"):
            self.assertIn(n, self.wb.defined_names)

    def test_tags_and_no_lb_mb_mix(self):
        # the build's own lint already ran with the tag index; check the
        # rows really are tagged by looking at the formulas' references
        measures = [r for r in self.registry if r.sheet == "Measures"]
        self.assertGreater(len(measures), 1000)
        tlb = set(c.row for k in KEYS for c in self._cells(
            self.wb, "M_%s_tLB" % k))
        tmb = set(c.row for k in KEYS for c in self._cells(
            self.wb, "M_%s_tMB" % k))
        self.assertEqual(len(tlb), 6)
        self.assertEqual(len(tmb), 6)
        for r in measures:
            self.assertEqual(lint_formula(r.text, "Measures"), [], r.text)
            if r.row + 1 in tlb:
                self.assertNotIn("EF_ResidualMB", r.text)
            if r.row + 1 in tmb:
                self.assertNotIn("EF_GridLB", r.text)

    def test_compare_cached_clean(self):
        bad, n = compare_cached(self.data)
        self.assertEqual(bad, [])
        self.assertGreater(n, 5000)

    def test_location_based_override(self):
        out = evaluate_names(
            self.data, ["M_RC_tSel", "M_LED_tSel", "M_EF_Sel"],
            overrides={("Inputs", "G9"): "Location-based"})
        lb = compute_measures(INPUTS, "Location-based", "Central")
        self.assertTrue(all(x == 0 for x in out["M_RC_tSel"]))
        for g, w in zip(out["M_LED_tSel"], lb.by_key("LED").t_lb):
            self.assertLessEqual(abs(g - w), REL * max(1.0, abs(w)))
        for g, w in zip(out["M_EF_Sel"], lb.f_lb):
            self.assertLessEqual(abs(g - w), REL * max(1.0, abs(w)))

    def test_scenario_override(self):
        out = evaluate_names(self.data, ["M_CarbonPrice"],
                             overrides={("Inputs", "G8"): "High"})
        hi = compute_measures(INPUTS, "Market-based", "High")
        for g, w in zip(out["M_CarbonPrice"], hi.carbon):
            self.assertLessEqual(abs(g - w), REL * max(1.0, abs(w)))

    def test_header_note(self):
        self.assertIn("Pre-tax, ex-VAT, real EUR 2025",
                      self.wb["Measures"]["B3"].value)


if __name__ == "__main__":
    unittest.main()
