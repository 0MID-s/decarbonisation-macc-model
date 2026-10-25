"""Disclosure sheet tests (plan 05-06)."""
import io
import re
import unittest

import openpyxl

import build
from config import disclosure as d
from config.security import host_allowed
from qa.formula_eval import compare_cached
from xl.sheets import disclosure as ds


class DisclosureSheet(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.data, _ = build.build_bytes("main", build.load_dataset("config"))
        cls.wb = openpyxl.load_workbook(io.BytesIO(cls.data))
        cls.ws = cls.wb["Disclosure"]
        cls.cells = [c for row in cls.ws.iter_rows() for c in row
                     if c.value is not None]
        hdr = [c for c in cls.cells if c.value == "Standard"][0]
        cls.hdr_row = hdr.row
        n = len(d.DISCLOSURE_ROWS)
        cls.rows = list(range(cls.hdr_row + 1, cls.hdr_row + 1 + n))

    def test_header_columns(self):
        got = [self.ws.cell(self.hdr_row, 2 + i).value
               for i in range(len(ds.HEADERS))]
        self.assertEqual(got, list(ds.HEADERS))

    def test_strength_cells_and_validation(self):
        for r in self.rows:
            self.assertIn(self.ws.cell(r, 8).value, d.STRENGTHS)
        dvs = self.ws.data_validations.dataValidation
        self.assertEqual(len(dvs), 1)
        dv = dvs[0]
        self.assertEqual(dv.type, "list")
        self.assertEqual(dv.formula1, '"Direct,Supports,Not covered"')
        self.assertEqual(str(dv.sqref), "H%d:H%d" % (self.rows[0],
                                                    self.rows[-1]))

    def test_key_text(self):
        text = " ".join(str(c.value) for c in self.cells)
        for s in ("Direct:", "Supports:", "Not covered:",
                  "representative firm built from sector averages"):
            self.assertIn(s, text)

    def test_not_covered_rows(self):
        n = sum(1 for r in self.rows
                if self.ws.cell(r, 8).value == "Not covered")
        self.assertGreaterEqual(n, 5)

    def test_only_name_formulas(self):
        pat = re.compile(r"^=[A-Za-z_][A-Za-z0-9_]*$")
        n = 0
        for c in self.cells:
            if isinstance(c.value, str) and c.value.startswith("="):
                n += 1
                self.assertRegex(c.value, pat)
        self.assertGreater(n, 10)

    def test_hyperlinks(self):
        n_int = n_ext = 0
        for c in self.cells:
            hl = c.hyperlink
            if hl is None:
                continue
            if hl.target:
                n_ext += 1
                self.assertTrue(host_allowed(hl.target), hl.target)
            else:
                n_int += 1
                sheet, _, addr = hl.location.partition("!")
                self.assertIn(sheet, self.wb.sheetnames)
                self.assertIsNotNone(self.wb[sheet][addr].value, hl.location)
        self.assertEqual(n_ext, len(d.STD_REFS))
        self.assertGreater(n_int, 20)

    def test_flag_visible(self):
        text = " ".join(str(c.value) for c in self.cells)
        self.assertIn("2026/1563", text)
        self.assertIn("EUR-Lex page not machine-readable", text)
        self.assertIn("Reference not verified", text)
        self.assertIn("amendment text not read", text)

    def test_dashboard_link(self):
        locs = {c.hyperlink.location for row in self.wb["Dashboard"].iter_rows()
                for c in row if c.hyperlink is not None
                and c.hyperlink.location}
        self.assertIn("Disclosure!A1", locs)

    def test_cached_values_match(self):
        bad, n = compare_cached(self.data)
        self.assertEqual(bad, [])
        self.assertGreater(n, 100)


if __name__ == "__main__":
    unittest.main()
