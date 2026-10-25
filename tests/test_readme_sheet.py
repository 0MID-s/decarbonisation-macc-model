"""ReadMe sheet tests (plan 05-05)."""
import io
import re
import unittest

import openpyxl

import build
from checkmodel.baseline_reference import model_values
from qa.formula_eval import evaluate_names
from xl.sheets import readme as rm


class ReadMeTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.ds = build.load_dataset("config")
        cls.data, _ = build.build_bytes("main", cls.ds)
        cls.wb = openpyxl.load_workbook(io.BytesIO(cls.data))
        cls.ws = cls.wb["ReadMe"]

    def _cells(self):
        return [c for row in self.ws.iter_rows() for c in row
                if c.value is not None]

    def test_cover_master_flag(self):
        cv = self.wb["Cover"]
        self.assertEqual(cv["F7"].value, "=Chk_Master")
        self.assertTrue(cv.conditional_formatting)
        self.assertEqual(cv["E7"].value, "Model checks")

    def test_second_tab(self):
        self.assertEqual(self.wb.sheetnames[1], "ReadMe")

    def test_sections_present(self):
        text = [str(c.value) for c in self._cells()]
        for s in ("How the model flows", "How to use the selectors",
                  "What never changes", "Trace this",
                  "Conventions in brief"):
            self.assertIn(s, text)

    def test_flow_table_links_every_present_sheet(self):
        linked = {c.value for c in self._cells()
                  if c.hyperlink is not None
                  and c.hyperlink.location.endswith("!A1")}
        for n in rm.FLOW_ORDER:
            if n in self.wb.sheetnames:
                self.assertIn(n, linked)

    def test_trace_links(self):
        starts = {t[0] for t in rm.TRACES}
        links = [c for c in self._cells()
                 if c.value in starts and c.hyperlink is not None]
        self.assertTrue(6 <= len(links) <= 8, len(links))
        for c in links:
            sheet, _, addr = c.hyperlink.location.partition("!")
            self.assertIsNotNone(self.wb[sheet][addr].value,
                                 c.hyperlink.location)

    def test_selector_links(self):
        locs = {c.hyperlink.location for c in self._cells()
                if c.hyperlink is not None}
        self.assertIn("Inputs!G8", locs)
        self.assertIn("Inputs!G9", locs)
        self.assertIn("Inputs!G10", locs)

    def test_prose_has_no_input_values(self):
        vals = set(model_values(self.ds.inputs).values())
        prose = " ".join(
            [t for t, _ in rm.SELECTORS] + list(rm.NEVER_CHANGES)
            + [rm.NPV_INVARIANT, rm.FLOW_ARROWS]
            + [t[2] for t in rm.TRACES] + list(rm.KEY_OUTPUTS.values()))
        prose = re.sub(r"(Phase|Scope|Step) \d", "", prose)
        for tok in re.findall(r"(?<![\w.-])\d+(?:\.\d+)?(?![\w.-])", prose):
            x = float(tok)
            if 1900 <= x <= 2100:
                continue
            self.assertNotIn(x, vals, tok)

    def test_npv_invariant_claim_holds(self):
        """The sentence on the sheet is only valid if M_Portfolio_NPV is the
        same in all 12 selector states."""
        from xl.sheets.readme import NPV_INVARIANT  # noqa: F401
        sel = {}
        holder = []
        build.build_bytes("main", self.ds,
                          extra_builders=[lambda wb, s, c: holder.append(c)])
        for n in ("Sel_Scenario", "Sel_Basis", "Sel_DirectETS"):
            sel[n] = holder[0].target_of(n)
        values = set()
        for sc in ("Low", "Central", "High"):
            for b in ("Location-based", "Market-based"):
                for e in ("Off", "On"):
                    ov = {sel["Sel_Scenario"]: sc, sel["Sel_Basis"]: b,
                          sel["Sel_DirectETS"]: e}
                    v = evaluate_names(self.data, ["M_Portfolio_NPV"],
                                       overrides=ov)["M_Portfolio_NPV"]
                    values.add(round(v, 6))
        self.assertEqual(len(values), 1, values)


if __name__ == "__main__":
    unittest.main()
