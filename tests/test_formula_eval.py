"""Formula evaluator tests (plan 02-14)."""
import io
import unittest

import openpyxl
import xlsxwriter

import build
from qa.formula_eval import (
    SUPPORTED_FUNCTIONS, FormulaEvalError, FormulaEvaluator, compare_cached,
    evaluate_names)


def _wb(fill):
    buf = io.BytesIO()
    wb = xlsxwriter.Workbook(buf, {"in_memory": True})
    fill(wb)
    wb.close()
    return openpyxl.load_workbook(io.BytesIO(buf.getvalue()))


def _basic(wb):
    s = wb.add_worksheet("Sheet1")
    o = wb.add_worksheet("Other")
    s.write("A1", 2)
    s.write("B1", 3)
    s.write_formula("C1", "=A1+B1")
    s.write_formula("D1", "=Other!$B$2")
    s.write_formula("E1", "=Cell_A1*2")
    s.write_formula("F1", "=SUM(Rng_L)")
    s.write_formula("G1", '=IF(1<2,"OK","ERROR")')
    s.write_formula("H1", "=IF(1<>2,1,0)+IF(2>=2,1,0)+IF(2<=1,1,0)")
    s.write_formula("I1", "=MATCH(5,Other!A1:A4,1)")
    s.write_formula("J1", "=MATCH(0,Other!A1:A4,1)")
    s.write_formula("K1", "=INDEX(Other!A1:A4,3)")
    s.write_formula("L1", "=CHOOSE(2,10,20,30)")
    s.write_formula("M1", "=FOO(1)")
    s.write_formula("N1", "=_xlfn.XLOOKUP(1,2,3)")
    s.write_formula("O1", "=P1")
    s.write_formula("P1", "=O1")
    s.write_formula("Q1", "=-2^2")
    s.write_formula("R1", "=A1__x")
    s.write_row("A2", [1, 2, 3])
    o.write("B2", 7)
    o.write_column("A1", [1, 4, 6, 9])
    wb.define_name("Cell_A1", "=Sheet1!$A$1")
    wb.define_name("Rng_L", "=Sheet1!$A$2:$C$2")


class EvaluatorUnit(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.wb = _wb(_basic)

    def ev(self, overrides=None):
        return FormulaEvaluator(self.wb, overrides)

    def test_basic_and_cross_sheet_and_names(self):
        e = self.ev()
        self.assertEqual(e.value("Sheet1", "C1"), 5)
        self.assertEqual(e.value("Sheet1", "D1"), 7)
        self.assertEqual(e.value("Sheet1", "E1"), 4)
        self.assertEqual(e.value("Sheet1", "F1"), 6)
        self.assertEqual(e.name_target("Rng_L"), ("Sheet1", "A2:C2"))

    def test_if_and_comparisons(self):
        e = self.ev()
        self.assertEqual(e.value("Sheet1", "G1"), "OK")
        self.assertEqual(e.value("Sheet1", "H1"), 2)

    def test_match_index_choose(self):
        e = self.ev()
        self.assertEqual(e.value("Sheet1", "I1"), 2)
        self.assertEqual(e.value("Sheet1", "K1"), 6)
        self.assertEqual(e.value("Sheet1", "L1"), 20)
        with self.assertRaises(FormulaEvalError):
            e.value("Sheet1", "J1")

    def test_overrides(self):
        e = self.ev({("Sheet1", "A1"): 10})
        self.assertEqual(e.value("Sheet1", "C1"), 13)
        # override keys are normalised like cell lookups (IN-07)
        e = self.ev({("Sheet1", "$a$1"): 10})
        self.assertEqual(e.value("Sheet1", "C1"), 13)
        e = self.ev({("Sheet1", "C1"): 99})
        self.assertEqual(e.value("Sheet1", "C1"), 99)

    def test_rejections(self):
        for coord in ("M1", "N1", "Q1", "R1"):
            with self.assertRaises(FormulaEvalError, msg=coord):
                self.ev().value("Sheet1", coord)

    def test_cycle(self):
        with self.assertRaises(FormulaEvalError):
            self.ev().value("Sheet1", "O1")

    def test_whitelist(self):
        self.assertIn("MATCH", SUPPORTED_FUNCTIONS)
        self.assertNotIn("EVAL", SUPPORTED_FUNCTIONS)


def _semantics(wb):
    s = wb.add_worksheet("Sheet1")
    s.write_row("B1", [1, 2])
    s.write_row("B2", [3, 4])
    s.write_string("D1", "LB")
    s.write_formula("A1", '="a"*3')
    s.write_formula("A2", '="2"*3')
    s.write_formula("A3", '="x"+"y"')
    s.write_formula("A4", "=B1:B2*2")
    s.write_formula("A5", '=IF("LB"="lb",1,0)')
    s.write_formula("A6", "=INDEX(B1:C2,3)")
    s.write_formula("A7", "=IF(1<2<3,1,0)")
    s.write_formula("A8", "=1=1=1")
    s.write_formula("A9", '=IF(D1="lb",1,0)+IF(D1<>"Lb",10,0)')
    s.write_formula("A10", '=IF("a"<"B",1,0)')
    s.write_formula("A11", '=IF(1="1",1,0)')
    s.write_formula("A12", "=B1:B2")
    s.write_formula("A13", "=-B1:B2")
    s.write_formula("A14", "=IF(AND(1<2,2<3),1,0)+IF((1<2)=TRUE,10,0)")
    s.write_formula("A15", "=INDEX(B1:B2,2)+INDEX(B2:C2,2)")
    s.write_formula("A16", "=TRUE+1")
    s.write_formula("A17", '=-"a"')
    s.write_formula("A18", '=IF("x",1,0)')


class EvaluatorSemantics(unittest.TestCase):
    """WR-01: refuse where Python and Excel semantics differ."""

    @classmethod
    def setUpClass(cls):
        cls.wb = _wb(_semantics)

    def v(self, coord):
        return FormulaEvaluator(self.wb).value("Sheet1", coord)

    def refuses(self, coord):
        with self.assertRaises(FormulaEvalError, msg=coord):
            self.v(coord)

    def test_string_arithmetic_refused(self):
        for coord in ("A1", "A2", "A3", "A17"):
            self.refuses(coord)

    def test_range_arithmetic_and_range_result_refused(self):
        for coord in ("A4", "A12", "A13"):
            self.refuses(coord)

    def test_text_comparison_case_insensitive(self):
        self.assertEqual(self.v("A5"), 1)
        self.assertEqual(self.v("A9"), 1)
        self.assertEqual(self.v("A10"), 1)

    def test_mixed_type_comparison_refused(self):
        self.refuses("A11")

    def test_index_2d_refused_1d_ok(self):
        self.refuses("A6")
        self.assertEqual(self.v("A15"), 7)

    def test_chained_comparison_refused(self):
        self.refuses("A7")
        self.refuses("A8")
        # one comparison per argument or parenthesis level is fine
        self.assertEqual(self.v("A14"), 11)

    def test_logical_arithmetic_and_text_condition(self):
        self.assertEqual(self.v("A16"), 2)
        self.refuses("A18")


class WorkbookFormulas(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        ds = build.load_dataset("config")
        cls.data, _ = build.build_bytes("main", ds)
        cls.fx = build.build_bytes("main", build.load_dataset("fixtures"))[0]

    def test_config_matches_cached(self):
        bad, n = compare_cached(self.data)
        self.assertGreaterEqual(n, 800)
        self.assertEqual(bad, [], "\n".join(bad[:10]))

    def test_fixtures_matches_cached(self):
        bad, n = compare_cached(self.fx)
        self.assertEqual(bad, [], "\n".join(bad[:10]))

    def test_mutation_changes_dependents(self):
        wf = openpyxl.load_workbook(io.BytesIO(self.data))
        sheet, ref = FormulaEvaluator(wf).name_target("Fin_FX_USDperEUR_2010")
        cell = ref.split(":")[0]
        want = openpyxl.load_workbook(io.BytesIO(self.data), data_only=True)[sheet][cell].value
        self.assertIsInstance(want, (int, float))
        # the override beats the cell's own formula: dependents must move
        bad, _ = compare_cached(self.data, overrides={(sheet, cell): want * 1.5})
        self.assertTrue(any("Carbon!" in b for b in bad), bad[:5])


class FirstYearRobustness(unittest.TestCase):
    """CR-01: Time_FirstYear is an editable input; Carbon paths must follow.

    With the pre-fix header-derived anchors, a 2027 run labelled the 2030 NGFS
    point as 2031 and the paths were silently wrong.
    """
    KEYS = ("CP", "NDC", "NZ50")

    @classmethod
    def setUpClass(cls):
        from checkmodel.carbon_reference import carbon_paths, interp_annual
        ds = build.load_dataset("config")
        data, _ = build.build_bytes("main", ds)
        cls.wf = openpyxl.load_workbook(io.BytesIO(data))
        cls.paths = carbon_paths(ds.inputs)
        cls.interp = staticmethod(interp_annual)
        sheet, ref = FormulaEvaluator(cls.wf).name_target("Time_FirstYear")
        cls.fy = (sheet, ref.split(":")[0])

    def _ev(self, year):
        return FormulaEvaluator(self.wf, {self.fy: year})

    def _cell(self, ev, name):
        sheet, ref = ev.name_target(name)
        return sheet, ref.split(":")[0]

    def test_first_year_shift(self):
        ev = self._ev(2027)
        self.assertEqual(ev.value("Carbon", "N5"), 2027)
        ws = self.wf["Carbon"]
        for key in self.KEYS:
            ays, avs, _ = self.paths[key]
            _, first = self._cell(ev, "Path_Carbon_" + key)
            row = ws[first].row
            yr_row = row - 6
            self.assertEqual(ws.cell(row=yr_row, column=2).value,
                             "NGFS anchor years")
            got_years = [ev.value("Carbon", ws.cell(row=yr_row, column=4 + i).coordinate)
                         for i in range(len(ays))]
            self.assertEqual(got_years, [2025, 2030, 2035, 2040, 2045, 2050, 2055])
            want = self.interp(ays, avs, list(range(2027, 2057)))
            c0 = ws[first].column
            for k, w in enumerate(want):
                g = ev.value("Carbon", ws.cell(row=row, column=c0 + k).coordinate)
                self.assertAlmostEqual(g / w, 1.0, delta=1e-12)
            s2, c = self._cell(ev, "Carb_Check_Anchors_" + key)
            self.assertEqual(ev.value(s2, c), "OK")

    def test_first_year_check_flags(self):
        ev = self._ev(2025)
        for key in self.KEYS:
            s2, c = self._cell(ev, "Carb_Check_Anchors_" + key)
            self.assertEqual(ev.value(s2, c), "ERROR")

    def test_edited_point_year_flags(self):
        """WR-02: a point year edited to 2026 must read ERROR, even though
        the years still ascend (2025 < 2026 < 2035)."""
        ws = self.wf["Carbon"]
        for key in self.KEYS:
            ev0 = FormulaEvaluator(self.wf)
            sheet, first = self._cell(ev0, "Path_Carbon_" + key)
            yr_row = ws[first].row - 6
            f = ws.cell(row=yr_row, column=5).value  # E: the 2030 point
            self.assertTrue(f.startswith("=Inputs!"), f)
            ref = f[len("=Inputs!"):].replace("$", "")
            self.assertEqual(ev0.value("Inputs", ref), 2030)
            ev = FormulaEvaluator(self.wf, {("Inputs", ref): 2026})
            s2, c = self._cell(ev, "Carb_Check_Anchors_" + key)
            self.assertEqual(ev.value(s2, c), "ERROR", key)

    def test_edited_y0_year_flags(self):
        ev0 = FormulaEvaluator(self.wf)
        y0 = self._cell(ev0, "Carbon_NGFS_Y0Year")
        ev = FormulaEvaluator(self.wf, {y0: 2020})
        for key in self.KEYS:
            s2, c = self._cell(ev, "Carb_Check_Anchors_" + key)
            self.assertEqual(ev.value(s2, c), "ERROR", key)


def _bytes(data, formulas, names=None):
    """Workbook bytes: data cells on Sheet1, formulas in column H (H1..)."""
    buf = io.BytesIO()
    wb = xlsxwriter.Workbook(buf, {"in_memory": True})
    ws = wb.add_worksheet("Sheet1")
    for coord, v in (data or {}).items():
        if isinstance(v, str):
            ws.write_string(coord, v)
        else:
            ws.write(coord, v)
    for i, f in enumerate(formulas):
        ws.write_formula("H%d" % (i + 1), f)
    for n, ref in (names or {}).items():
        wb.define_name(n, ref)
    wb.close()
    return buf.getvalue()


def _calc(formula, data=None, overrides=None):
    wf = openpyxl.load_workbook(io.BytesIO(_bytes(data, [formula])))
    return FormulaEvaluator(wf, overrides).value("Sheet1", "H1")


ROW = {"A1": 1, "B1": 2, "C1": 3}
POS = {"A1": 10, "B1": 20, "C1": 30}


class TestPhase3Functions(unittest.TestCase):
    def refuses(self, formula, data=None):
        with self.assertRaises(FormulaEvalError, msg=formula):
            _calc(formula, data)

    def test_power(self):
        self.assertEqual(_calc("=2^3"), 8)
        self.assertAlmostEqual(_calc("=(1+0.1)^2"), 1.21, places=12)
        self.assertAlmostEqual(_calc("=2^-2"), 0.25, places=12)
        self.assertEqual(_calc("=2*3^2"), 18)

    def test_power_refusals(self):
        self.refuses("=-2^2")
        self.refuses("=2^3^2")  # Excel left-assoc, Python right-assoc
        self.refuses("=0^-1")   # #DIV/0! error value reaches the cell
        self.refuses("=(-8)^0.5")  # #NUM!

    def test_concat(self):
        self.assertEqual(_calc('="a"&"b"'), "ab")
        self.assertEqual(_calc('="x: "&A1', {"A1": "Central"}), "x: Central")
        self.assertEqual(_calc('=IF("a"&"b"="AB",1,0)'), 1)
        self.refuses('="a"&1')
        self.refuses('=1&2')

    def test_sumproduct(self):
        self.assertEqual(_calc("=SUMPRODUCT(A1:C1,A2:C2)",
                               {"A1": 1, "B1": 2, "C1": 3, "A2": 4, "B2": 5, "C2": 6}), 32)
        self.refuses("=SUMPRODUCT(A1:C1,A2:B2)", {"A1": 1, "B1": 2, "C1": 3, "A2": 4, "B2": 5})

    def test_npv(self):
        self.assertAlmostEqual(_calc("=NPV(0.1,110,121)"), 200.0, places=9)
        self.assertAlmostEqual(_calc("=NPV(0.1,A1:C1)", {"A1": 110, "B1": 121, "C1": 0}), 200.0, places=9)

    def test_irr(self):
        self.assertAlmostEqual(_calc("=IRR(A1:B1)", {"A1": -100, "B1": 110}), 0.1, places=9)
        # guess is ignored
        self.assertAlmostEqual(_calc("=IRR(A1:B1,0.5)", {"A1": -100, "B1": 110}), 0.1, places=9)
        # replacement-capex shape: three sign changes, single root
        flows = {"A1": -100, "B1": 60, "C1": 60, "D1": -10, "E1": 60}
        r = _calc("=IRR(A1:E1)", flows)
        npv = sum(v / (1 + r) ** i for i, v in enumerate([-100, 60, 60, -10, 60]))
        self.assertAlmostEqual(npv, 0.0, places=8)

    def test_irr_refusals(self):
        # two roots at 10% and 20%
        self.refuses("=IRR(A1:C1)", {"A1": -100, "B1": 230, "C1": -132})
        self.refuses("=IRR(A1:C1)", POS)
        self.refuses("=IRR(A1:C1)+1", POS)

    def test_mod_round(self):
        self.assertEqual(_calc("=MOD(29,15)"), 14)
        self.assertEqual(_calc("=MOD(29,8)"), 5)
        self.assertEqual(_calc("=MOD(-1,8)"), 7)
        self.refuses("=MOD(5,0)")
        self.assertEqual(_calc("=ROUND(2.5,0)"), 3)
        self.assertEqual(_calc("=ROUND(-2.5,0)"), -3)
        self.assertEqual(_calc("=ROUND(2.4,0)"), 2)
        self.refuses("=ROUND(1.2345,2)")

    def test_rank(self):
        d = {"A1": 3, "B1": 5, "C1": 1}
        self.assertEqual(_calc("=RANK(5,A1:C1,1)", d), 3)
        self.assertEqual(_calc("=RANK(5,A1:C1,0)", d), 1)
        self.assertEqual(_calc("=RANK(5,A1:C1)", d), 1)
        self.assertEqual(_calc("=RANK(3,A1:D1,1)", dict(d, D1="text")), 2)
        self.refuses('=RANK("a",A1:C1,1)', d)
        self.assertEqual(_calc('=IFERROR(RANK("a",A1:C1,1),-1)', d), -1)

    def test_countif_count(self):
        d = {"A1": 5, "B1": 5, "C1": 0, "D1": "x"}
        self.assertEqual(_calc("=COUNTIF(A1:D1,5)", d), 2)
        self.assertEqual(_calc("=COUNTIF(A1:D1,0)", d), 1)
        self.assertEqual(_calc('=COUNTIF(A1:D1,"X")', d), 1)
        self.refuses('=COUNTIF(A1:D1,">1")', d)
        self.assertEqual(_calc("=COUNT(A1:D1)", d), 3)

    def test_logic(self):
        self.assertEqual(_calc("=OR(1>2,2>1)"), True)
        self.assertEqual(_calc("=OR(1>2,2>3)"), False)
        self.assertEqual(_calc("=NOT(1>2)"), True)
        self.assertEqual(_calc("=NOT(0)"), True)
        self.assertEqual(_calc("=ISNUMBER(3)"), True)
        self.assertEqual(_calc("=ISNUMBER(A1)", {"A1": "t"}), False)
        self.assertEqual(_calc("=ISNUMBER(TRUE)"), False)

    def test_match_type0(self):
        names = {"A1": "Low", "B1": "Central", "C1": "High"}
        self.assertEqual(_calc('=MATCH("central",A1:C1,0)', names), 2)
        self.refuses('=MATCH("none",A1:C1,0)', names)
        self.refuses("=IF(MATCH(9,A1:C1,0)>0,1,0)", POS)
        self.assertEqual(_calc("=MATCH(20,A1:C1,0)", POS), 2)

    def test_lazy_errors(self):
        self.assertEqual(_calc('=IF(1>0,"n/a",IRR(A1:C1))', POS), "n/a")
        self.assertEqual(_calc('=IFERROR(IRR(A1:C1),"none")', POS), "none")
        self.assertEqual(_calc("=IF(1>0,1,1/0)"), 1)
        self.assertEqual(_calc("=IFERROR(1/0,7)"), 7)
        self.assertEqual(_calc("=IFERROR(3,7)"), 3)
        self.refuses("=1/0")

    def test_whitelist_and_injection(self):
        for n in ("SUMPRODUCT", "NPV", "IRR", "MOD", "ROUND", "RANK", "COUNTIF",
                  "COUNT", "OR", "NOT", "IFERROR", "ISNUMBER"):
            self.assertIn(n, SUPPORTED_FUNCTIONS)
        self.refuses("=__import__(1)")
        self.refuses("=EVAL(1)")


class TestEvaluateNames(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        data = {"B1": "Central", "A5": 1, "A6": 2}
        formulas = [
            '=IF(B1="High",100,10)',
            '=IF(B1="High",100,10)+1',
        ]
        buf = io.BytesIO()
        wb = xlsxwriter.Workbook(buf, {"in_memory": True})
        ws = wb.add_worksheet("Sheet1")
        ws.write_string("B1", "Central")
        for j, c in enumerate("CDEF"):
            ws.write_formula("%s2" % c, '=IF($B$1="High",%d,%d)' % (100 * (j + 1), j + 1))
            ws.write_formula("%s3" % c, "=%s2*2" % c)
        ws.write_formula("A4", '=IF(B1="High",100,10)')
        wb.define_name("Name_A", "=Sheet1!$A$4")
        wb.define_name("Row_A", "=Sheet1!$C$2:$F$2")
        wb.define_name("Grid_A", "=Sheet1!$C$2:$D$3")
        wb.close()
        cls.data = buf.getvalue()

    def test_scalar_name_with_override(self):
        self.assertEqual(evaluate_names(self.data, ["Name_A"])["Name_A"], 10)
        out = evaluate_names(self.data, ["Name_A"], overrides={("Sheet1", "B1"): "High"})
        self.assertEqual(out["Name_A"], 100)

    def test_row_range_name(self):
        self.assertEqual(evaluate_names(self.data, ["Row_A"])["Row_A"], [1, 2, 3, 4])
        out = evaluate_names(self.data, ["Row_A"], overrides={("Sheet1", "$B$1"): "High"})
        self.assertEqual(out["Row_A"], [100, 200, 300, 400])

    def test_two_d_range_row_major(self):
        out = evaluate_names(self.data, ["Grid_A"])
        self.assertEqual(out["Grid_A"], [1, 2, 2, 4])

    def test_unknown_name_refused(self):
        with self.assertRaises(FormulaEvalError):
            evaluate_names(self.data, ["No_Such_Name"])
